from __future__ import annotations

import copy
import shutil
import subprocess
import tempfile
import unicodedata
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from lxml import etree


ROOT = Path(__file__).resolve().parents[1]
INPUT_DOCX = ROOT / "output" / "doc" / "第二部分_材料与方法_学术论文风格重写稿_v2_公式与三线表.docx"
OUTPUT_DOCX = ROOT / "output" / "doc" / "第二部分_材料与方法_学术论文风格重写稿_v2_公式编号修订版.docx"

MATH_NS = {"m": "http://schemas.openxmlformats.org/officeDocument/2006/math"}


CORRECTED_FORMULAS = {
    10: (
        r"C_{ij}="
        r"\frac{\sum_{b=1}^{B} z^{A}_{b,i}z^{B}_{b,j}}"
        r"{\sqrt{\sum_{b=1}^{B}(z^{A}_{b,i})^{2}}\sqrt{\sum_{b=1}^{B}(z^{B}_{b,j})^{2}}},"
        r"\quad C\in\mathbb{R}^{d\times d}"
    ),
    11: (
        r"\mathcal{L}_{BT}="
        r"\sum_{i=1}^{d}(1-C_{ii})^{2}+"
        r"\lambda_{BT}\sum_{i\ne j}\left(C_{ij}\right)^{2}"
    ),
    14: (
        r"A^{\prime}_{n,j}="
        r"\frac{|A_{n,j}|}{\sum_{k=1}^{J}|A_{n,k}|+\epsilon},"
        r"\quad j=1,\ldots,J"
    ),
}

def latex_to_omml_para(latex: str, work_dir: Path) -> etree._Element:
    md_path = work_dir / "formula.md"
    docx_path = work_dir / "formula.docx"
    md_path.write_text(f"$${latex}$$\n", encoding="utf-8")
    subprocess.run(
        ["pandoc", str(md_path), "-o", str(docx_path)],
        check=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
    )
    xml = ZipFile(docx_path).read("word/document.xml")
    root = etree.fromstring(xml)
    omml = root.xpath("//m:oMathPara", namespaces=MATH_NS)
    if not omml:
        raise RuntimeError(f"Pandoc did not emit an OMML formula for: {latex}")
    element = copy.deepcopy(omml[0])
    for node in element.xpath(".//m:t", namespaces=MATH_NS):
        if node.text:
            node.text = "".join(
                ch for ch in node.text if unicodedata.category(ch) != "Cf"
            )
    return element


def paragraph_has_math(paragraph) -> bool:
    return bool(paragraph._p.xpath(".//*[local-name()='oMathPara' or local-name()='oMath']"))


def replace_paragraph_formula(paragraph, omml_para: etree._Element) -> None:
    p = paragraph._p
    for child in list(p):
        if child.tag != qn("w:pPr"):
            p.remove(child)
    p.append(omml_para)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(4)


def text_width_twips(doc: Document) -> int:
    section = doc.sections[0]
    return int(section.page_width.twips - section.left_margin.twips - section.right_margin.twips)


def get_inline_math(paragraph) -> etree._Element:
    math_nodes = paragraph._p.xpath(".//*[local-name()='oMath']")
    if not math_nodes:
        raise RuntimeError("Formula paragraph does not contain an inline oMath node")
    return copy.deepcopy(math_nodes[0])


def ensure_ppr(paragraph) -> OxmlElement:
    p_pr = paragraph._p.pPr
    if p_pr is None:
        p_pr = OxmlElement("w:pPr")
        paragraph._p.insert(0, p_pr)
    return p_pr


def set_equation_tabs(paragraph, text_width: int) -> None:
    p_pr = ensure_ppr(paragraph)
    old_tabs = p_pr.find(qn("w:tabs"))
    if old_tabs is not None:
        p_pr.remove(old_tabs)
    tabs = OxmlElement("w:tabs")
    center = OxmlElement("w:tab")
    center.set(qn("w:val"), "center")
    center.set(qn("w:pos"), str(text_width // 2))
    tabs.append(center)
    right = OxmlElement("w:tab")
    right.set(qn("w:val"), "right")
    right.set(qn("w:pos"), str(text_width))
    tabs.append(right)
    p_pr.append(tabs)


def make_tab_run() -> OxmlElement:
    run = OxmlElement("w:r")
    tab = OxmlElement("w:tab")
    run.append(tab)
    return run


def make_number_run(number: int) -> OxmlElement:
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")
    fonts.set(qn("w:eastAsia"), "宋体")
    r_pr.append(fonts)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "21")
    r_pr.append(size)
    run.append(r_pr)
    text = OxmlElement("w:t")
    text.text = f"({number})"
    run.append(text)
    return run


def replace_formula_paragraph_with_numbered_tabs(paragraph, number: int, text_width: int) -> None:
    inline_math = get_inline_math(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(4)
    set_equation_tabs(paragraph, text_width)
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    paragraph._p.append(make_tab_run())
    paragraph._p.append(inline_math)
    paragraph._p.append(make_tab_run())
    paragraph._p.append(make_number_run(number))


def main() -> None:
    if not INPUT_DOCX.exists():
        raise FileNotFoundError(INPUT_DOCX)
    if shutil.which("pandoc") is None:
        raise RuntimeError("pandoc is required for LaTeX-to-Word-equation conversion")

    doc = Document(INPUT_DOCX)
    formula_paragraphs = [p for p in doc.paragraphs if paragraph_has_math(p)]
    if len(formula_paragraphs) != 15:
        raise RuntimeError(f"Expected 15 formula paragraphs, found {len(formula_paragraphs)}")

    with tempfile.TemporaryDirectory() as tmp:
        tmp_path = Path(tmp)
        for formula_number, latex in CORRECTED_FORMULAS.items():
            replace_paragraph_formula(
                formula_paragraphs[formula_number - 1],
                latex_to_omml_para(latex, tmp_path),
            )

    # Refresh after formula replacement, then convert every display formula to a
    # single paragraph with center and right tab stops: formula centered, number right.
    formula_paragraphs = [p for p in doc.paragraphs if paragraph_has_math(p)]
    width = text_width_twips(doc)
    for idx, paragraph in enumerate(formula_paragraphs, start=1):
        replace_formula_paragraph_with_numbered_tabs(paragraph, idx, width)

    OUTPUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_DOCX)
    print(f"numbered_formulas={len(formula_paragraphs)}")
    print(f"saved={OUTPUT_DOCX}")


if __name__ == "__main__":
    main()
