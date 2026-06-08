from __future__ import annotations

import csv
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE_WORKBOOKS = [
    Path(r"F:\CJZFile\EEG_M1\19例患者脑电数据完整性检查.xlsx"),
    Path(r"F:\CJZFile\EEG_M1\M1组病历记录表.xlsx"),
    Path(r"F:\CJZFile\EEG_M1\脑卒中患者信息记录表.xlsx"),
]
OUTPUT_CSV = ROOT / "results" / "tables" / "source_workbook_author_metadata_audit.csv"
OUTPUT_MD = ROOT / "docs" / "source_workbook_author_metadata_audit.md"


@dataclass(frozen=True)
class MetadataField:
    field_id: str
    required_evidence: str
    keywords: tuple[str, ...]
    current_interpretation: str
    author_action: str


FIELDS = [
    MetadataField(
        "ethics_approval",
        "ethics committee name, approval number, approval date, and applicable site",
        ("伦理", "伦理委员会", "审批", "批准", "IRB", "ethics", "approval"),
        "No source-workbook keyword hit can provide a submission-ready ethics statement.",
        "Provide the ethics approval document or author-approved ethics statement.",
    ),
    MetadataField(
        "informed_consent",
        "consent route, consent provider, written or waived consent, and data-sharing scope",
        ("知情", "同意", "consent", "waiver", "授权"),
        "No consent-form text or consent scope is available from the source workbooks.",
        "Provide consent wording, waiver route, and whether data sharing is covered.",
    ),
    MetadataField(
        "trial_or_study_registration",
        "registry name and number, or reason for non-registration",
        ("注册", "ChiCTR", "NCT", "clinicaltrials", "registry", "登记"),
        "No registry identifier is available from the source workbooks.",
        "Provide registry information or an author-approved non-registration statement.",
    ),
    MetadataField(
        "study_site_dates_design",
        "hospital/department, recruitment dates, and design classification",
        ("医院", "科", "病区", "招募", "入组", "日期", "时间", "门诊", "住院"),
        "Source workbooks contain limited timing-like clinical fields but not recruitment dates, site, or final design wording.",
        "Provide hospital/department, recruitment start/end dates, and prospective/retrospective design.",
    ),
    MetadataField(
        "eligibility_stroke_timing",
        "inclusion/exclusion criteria, stroke subtype, lesion criteria, and timing from stroke to EEG/tACS",
        ("纳入", "排除", "标准", "诊断", "缺血", "出血", "卒中", "脑卒中", "病程", "MMSE"),
        "Clinical workbooks contain disease-duration and MMSE-like fields, but not protocol-level eligibility or stroke-subtype criteria.",
        "Provide inclusion/exclusion criteria, subtype criteria, lesion/timing rules, and define what disease duration measures.",
    ),
    MetadataField(
        "tacs_device_electrodes",
        "device model, electrode dimensions/materials, and tACS protocol confirmation",
        ("tACS", "刺激", "电极", "设备", "仪器", "C3", "C4", "M1", "20Hz", "20 Hz"),
        "Source workbooks do not contain device model or electrode dimensions.",
        "Provide device model, electrode size/materials, and confirm target/frequency/intensity/session protocol.",
    ),
    MetadataField(
        "concurrent_rehabilitation",
        "whether conventional rehabilitation was delivered, with dose and content",
        ("康复", "训练", "作业治疗", "物理治疗", "常规", "PT", "OT", "rehabilitation"),
        "No source-workbook field verifies concurrent conventional rehabilitation dose or content.",
        "Confirm whether patients received tACS alone or tACS plus standardized rehabilitation, including dose/content.",
    ),
    MetadataField(
        "tacs_safety_adverse_events",
        "adverse events, tolerability, withdrawals, and monitoring method",
        ("脱落", "缺少数据", "不适", "难受", "高血压", "头痛", "依从", "不良", "安全", "退出"),
        "Source workbooks contain missing-data/dropout notes, including discomfort-related terms, but not a formal safety-monitoring dataset.",
        "Provide an adverse-event/tolerability summary or explicitly state that formal safety data were unavailable.",
    ),
    MetadataField(
        "fma_assessors_timing",
        "assessor credentials, blinding, and baseline/post-treatment assessment timing",
        ("FMA", "MBI", "BBT", "评估", "量表", "盲", "评分"),
        "Source workbooks contain clinical scale columns but not assessor credentials or blinding.",
        "Provide assessor training/credentials, blinding status, and assessment timing language.",
    ),
    MetadataField(
        "eeg_hardware_reference_impedance",
        "amplifier, acquisition software, cap system, original channels, reference, ground, and impedance",
        ("EEG", "脑电", "采集", "放大器", "帽", "导联", "参考", "地", "阻抗", "reference", "ground", "impedance"),
        "Workbook fields do not verify amplifier, cap, reference, ground, or impedance protocol.",
        "Provide acquisition hardware, cap/montage, reference/ground, impedance, and software details.",
    ),
    MetadataField(
        "resting_state_instructions",
        "eyes-open/eyes-closed order, target duration, fixation instruction, and drowsiness monitoring",
        ("睁眼", "闭眼", "静息", "眼", "fixation", "drowsiness", "困倦"),
        "The completeness workbook records EO/EC task availability but not acquisition instructions.",
        "Provide EO/EC order, fixation or eye instruction, target duration, and drowsiness monitoring.",
    ),
    MetadataField(
        "raw_eeg_preprocessing",
        "raw filtering, notch, re-reference, artifact handling, bad-channel handling, and export rules",
        ("预处理", "滤波", "陷波", "重参考", "参考", "ICA", "伪迹", "坏道", "插值", "artifact", "filter"),
        "Source workbooks do not document upstream raw EEG preprocessing before the provided EEGLAB files.",
        "Provide raw preprocessing protocol before .set/.fdt export.",
    ),
    MetadataField(
        "data_repository_doi_scope",
        "repository, DOI/accession, version, license, public/restricted scope, and request route",
        ("DOI", "仓储", "共享", "数据", "license", "repository", "Zenodo", "OSF"),
        "No final repository DOI, licence, or controlled-access route is present in the source workbooks.",
        "Provide repository DOI/accession, public data scope, restricted data scope, licence, and access route.",
    ),
    MetadataField(
        "funding_competing_acknowledgements",
        "funding, competing interests, author acknowledgements",
        ("基金", "资助", "经费", "利益冲突", "致谢", "funding", "competing"),
        "No funding, competing-interest, or acknowledgement evidence is present in the source workbooks.",
        "Provide author-approved funding, competing-interest, and acknowledgement statements.",
    ),
]


def main() -> None:
    OUTPUT_CSV.parent.mkdir(parents=True, exist_ok=True)
    rows = audit_workbooks()
    write_csv(rows)
    write_markdown(rows)
    print(OUTPUT_CSV)
    print(OUTPUT_MD)


def audit_workbooks() -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    workbook_summaries = inspect_workbooks()
    for field in FIELDS:
        refs: list[str] = []
        matched_terms: Counter[str] = Counter()
        workbook_hits: Counter[str] = Counter()
        for summary in workbook_summaries:
            for hit in summary["hits"]:
                if hit["term"] in field.keywords:
                    matched_terms[hit["term"]] += 1
                    workbook_hits[summary["workbook"]] += 1
                    if len(refs) < 8:
                        refs.append(f"{summary['workbook']}:{hit['sheet']}!{hit['cell']}")

        hit_count = sum(workbook_hits.values())
        evidence_grade = grade_field(field.field_id, hit_count)
        rows.append(
            {
                "field_id": field.field_id,
                "required_evidence": field.required_evidence,
                "source_workbooks_inspected": "; ".join(summary["workbook"] for summary in workbook_summaries),
                "keyword_hit_count": str(hit_count),
                "matched_terms": "; ".join(f"{term} ({count})" for term, count in matched_terms.most_common()) or "none",
                "nonidentifying_cell_refs": "; ".join(refs) or "none",
                "evidence_grade": evidence_grade,
                "current_interpretation": field.current_interpretation,
                "author_action": field.author_action,
            }
        )
    return rows


def inspect_workbooks() -> list[dict[str, object]]:
    summaries: list[dict[str, object]] = []
    all_terms = {term for field in FIELDS for term in field.keywords}
    for path in SOURCE_WORKBOOKS:
        if not path.exists():
            summaries.append({"workbook": path.name, "hits": []})
            continue
        wb = load_workbook(path, data_only=True, read_only=True)
        hits: list[dict[str, str]] = []
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    text = str(cell.value)
                    for term in all_terms:
                        if term.lower() in text.lower():
                            hits.append({"sheet": ws.title, "cell": cell.coordinate, "term": term})
        summaries.append({"workbook": path.name, "hits": hits})
    return summaries


def grade_field(field_id: str, hit_count: int) -> str:
    partial = {
        "eligibility_stroke_timing",
        "tacs_safety_adverse_events",
        "fma_assessors_timing",
        "resting_state_instructions",
    }
    if field_id in partial and hit_count > 0:
        return "partial_source_field_present_not_submission_ready"
    if hit_count > 0:
        return "keyword_hits_not_submission_ready"
    return "no_source_workbook_evidence"


def write_csv(rows: list[dict[str, str]]) -> None:
    with OUTPUT_CSV.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)


def write_markdown(rows: list[dict[str, str]]) -> None:
    counts = Counter(row["evidence_grade"] for row in rows)
    lines = [
        "# Source Workbook Author-Metadata Audit",
        "",
        "This audit scans the three current M1 source workbooks for non-identifying evidence relevant to author-supplied submission metadata. It reports keyword hits and workbook cell references, but does not expose patient names or individual source values. A hit is treated as evidence only when it can support a submission-ready protocol statement.",
        "",
        "## Summary",
        "",
    ]
    for grade, count in sorted(counts.items()):
        lines.append(f"- {grade}: {count}")
    lines.extend(
        [
            "",
            "## Field Audit",
            "",
            "| Field | Evidence grade | Keyword hits | Matched terms | Interpretation | Author action |",
            "|---|---|---:|---|---|---|",
        ]
    )
    for row in rows:
        lines.append(
            "| {field_id} | {evidence_grade} | {keyword_hit_count} | {matched_terms} | {current_interpretation} | {author_action} |".format(
                **{key: value.replace("|", "/") for key, value in row.items()}
            )
        )
    lines.extend(
        [
            "",
            "## Interpretation",
            "",
            "- Workbook fields can support selected descriptive statements, such as disease-duration availability, clinical-scale availability, EO/EC task completeness, and missing-data/dropout note categories.",
            "- The source workbooks do not contain submission-ready ethics, consent, registration, EEG acquisition hardware, upstream raw preprocessing, conventional-rehabilitation dose, repository DOI/licence, funding, competing-interest, or authorship evidence.",
            "- The final manuscript should keep these items as author-confirmed fields unless the author supplies protocol, ethics, device, rehabilitation, or repository documents.",
        ]
    )
    OUTPUT_MD.write_text("\n".join(lines) + "\n", encoding="utf-8")


if __name__ == "__main__":
    main()
