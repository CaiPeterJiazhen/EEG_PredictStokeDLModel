from __future__ import annotations

import csv
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "output" / "doc"
OUTPUT_DOCX = OUTPUT_DIR / "ResidualAware_SSL_CNN_Supplementary_Information.docx"

SUPPLEMENTARY_FIGURES = [
    (
        "Supplementary Figure 1",
        ROOT / "results" / "figures" / "nature" / "supplementary_error_subjects.png",
        "Repeated-error subject analysis. Post-hoc subject-level error summaries identify patients repeatedly misclassified across model variants or seeds. This analysis is exploratory and should be used to guide future cohort review rather than label revision.",
    ),
    (
        "Supplementary Figure 2",
        ROOT
        / "results"
        / "figures"
        / "explainability"
        / "mne_wpli_connectivity"
        / "mne_wpli_connectivity_contact_sheet.png",
        "MNE-rendered WPLI connectivity attribution maps. The top 20 attribution edges are shown for each EEG state and frequency band using the fixed 62-channel scalp layout. Red and blue indicate attribution sign, and line width scales with mean absolute attribution.",
    ),
    (
        "Supplementary Figure 3",
        ROOT / "results" / "figures" / "nature" / "supplementary_performance_precision.png",
        "Performance precision and validation-boundary audit. Panel a shows final-model bootstrap intervals for primary performance metrics. Panel b shows paired bootstrap differences for the final model relative to the no-SSL CNN and logistic EEG baseline; intervals crossing zero indicate that superiority was not established. For Brier score, negative differences favor the final model. Panel c shows the one-case movement in accuracy, sensitivity, and specificity for the 19-patient LOSO cohort.",
    ),
    (
        "Supplementary Figure 4",
        ROOT / "results" / "figures" / "nature" / "supplementary_clinical_incremental_value.png",
        "Exploratory clinical baseline and EEG incremental-value audit. Panel a compares selected clinical-only, EEG-plus-clinical, and EEG-only model metrics. Panel b shows paired bootstrap differences for EEG-plus-clinical candidates relative to the clinical-only logistic model. Panel c summarizes directional point estimates across all exploratory candidates after transforming each metric so positive values favor the candidate. Panel d plots ROC-AUC benefit against Brier-score benefit relative to the clinical-only model.",
    ),
]


SUPPLEMENTARY_TABLES = [
    (
        "Supplementary Table 1",
        "Subject-level model confidence intervals and calibration.",
        ROOT / "results" / "statistics" / "model_metric_confidence_intervals.csv",
        [
            "model_name",
            "accuracy",
            "balanced_accuracy",
            "roc_auc",
            "pr_auc",
            "brier_score",
            "ece",
            "calibration_intercept",
            "calibration_slope",
            "binomial_accuracy_p",
        ],
    ),
    (
        "Supplementary Table 2",
        "Paired model comparisons using subject-level bootstrap or McNemar tests.",
        ROOT / "results" / "statistics" / "model_pairwise_comparisons.csv",
        [
            "model_a",
            "model_b",
            "metric",
            "metric_a",
            "metric_b",
            "difference",
            "ci_low",
            "ci_high",
            "p_value_two_sided",
            "p_value",
        ],
    ),
    (
        "Supplementary Table 3",
        "Subject-level label permutation tests.",
        ROOT / "results" / "statistics" / "model_permutation_tests.csv",
        ["model_name", "metric", "observed", "p_value", "n_permutations", "n_valid_permutations"],
    ),
    (
        "Supplementary Table 4",
        "Exploratory clinical and EEG-plus-clinical paired bootstrap comparisons.",
        ROOT / "results" / "statistics" / "clinical_incremental_paired_bootstrap_comparison.csv",
        [
            "reference_model",
            "candidate_model",
            "metric",
            "metric_a",
            "metric_b",
            "difference",
            "ci_low",
            "ci_high",
            "p_value_two_sided",
        ],
    ),
    (
        "Supplementary Table 5",
        "Seed-level stability summary.",
        ROOT / "results" / "tables" / "seed_stability_table.csv",
        [
            "model",
            "mean_accuracy",
            "std_accuracy",
            "min_accuracy",
            "mean_roc_auc",
            "mean_pr_auc",
            "mean_brier",
        ],
    ),
    (
        "Supplementary Table 6",
        "Non-identifying EEG recording metadata summary.",
        ROOT / "results" / "tables" / "eeg_recording_summary.csv",
        [
            "group",
            "stage",
            "state",
            "n_records",
            "n_subjects",
            "srate_values_hz",
            "nbchan_values",
            "trials_values",
            "duration_sec_mean",
            "duration_sec_min",
            "duration_sec_max",
        ],
    ),
    (
        "Supplementary Table 7",
        "Participant-flow and safety source-note categories from de-identified project records.",
        ROOT / "results" / "tables" / "participant_flow_safety_source_notes.csv",
        [
            "row_type",
            "category",
            "n",
            "denominator",
            "manuscript_use",
            "notes",
        ],
    ),
    (
        "Supplementary Table 8",
        "PROBAST/TRIPOD+AI-oriented risk-of-bias and applicability audit.",
        ROOT / "results" / "tables" / "probast_tripod_ai_risk_audit.csv",
        [
            "domain",
            "assessment_item",
            "risk_of_bias",
            "applicability_concern",
            "mitigation_or_required_action",
        ],
    ),
    (
        "Supplementary Table 9",
        "Performance precision and validation-boundary audit for the 19-patient LOSO analysis.",
        ROOT / "results" / "tables" / "performance_precision_audit.csv",
        [
            "category",
            "item",
            "observed_value",
            "uncertainty_or_resolution",
            "interpretation",
            "evidence_source",
        ],
    ),
    (
        "Supplementary Table 10",
        "Claim-strength audit mapping central manuscript claims to evidence, acceptable wording, and overclaims to avoid.",
        ROOT / "results" / "tables" / "claim_strength_audit.csv",
        [
            "claim_id",
            "manuscript_section",
            "manuscript_claim",
            "evidence_strength",
            "allowed_wording",
            "overclaim_to_avoid",
            "verification_status",
        ],
    ),
    (
        "Supplementary Table 11",
        "AI model reporting card summarizing intended use, validation safeguards, reproducibility actions, and unsupported uses.",
        ROOT / "results" / "tables" / "model_reporting_card.csv",
        [
            "card_section",
            "reporting_item",
            "current_value",
            "reviewer_risk",
            "manuscript_action",
        ],
    ),
]


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def compact_value(value: str | None) -> str:
    if value is None:
        return ""
    text = str(value)
    if text == "" or text.lower() == "nan":
        return ""
    try:
        number = float(text)
    except ValueError:
        return shorten_model_name(text)
    if abs(number) >= 100:
        return f"{number:.0f}"
    if abs(number) >= 10:
        return f"{number:.2f}"
    return f"{number:.3f}"


def shorten_model_name(value: str) -> str:
    replacements = {
        "ML_EEG_updated_no_selector_logistic_l1": "Logistic EEG",
        "no_SSL_CNN_updated_sub05_sub28_seedensemble10": "CNN",
        "residual_aware_SSL_CNN_seedmean10": "Residual-aware SSL-CNN",
        "clinical_only_logistic_l2": "Clinical logistic L2",
        "eeg_clinical_logistic_l1_selectk100": "EEG+clinical logistic L1",
    }
    return replacements.get(value, value)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def write_cell(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(6.5)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(9)
    normal.paragraph_format.space_after = Pt(3)
    for style_name, size in [("Title", 17), ("Heading 1", 13), ("Heading 2", 10)]:
        style = styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True


def add_table(document: Document, label: str, caption: str, path: Path, columns: list[str]) -> None:
    rows = read_csv(path)
    document.add_heading(label, level=2)
    caption_p = document.add_paragraph()
    caption_run = caption_p.add_run(caption)
    caption_run.bold = True
    caption_run.font.size = Pt(8.5)

    table = document.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for index, column in enumerate(columns):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, "D9E2F3")
        write_cell(cell, column, bold=True)

    for row in rows:
        cells = table.add_row().cells
        for index, column in enumerate(columns):
            write_cell(cells[index], compact_value(row.get(column, "")))

    note = document.add_paragraph()
    note.add_run(f"Source: {path.relative_to(ROOT).as_posix()}").italic = True


def add_figure(document: Document, label: str, path: Path, caption: str, width: float) -> None:
    document.add_heading(label, level=2)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(path), width=Inches(width))
    caption_p = document.add_paragraph()
    caption_run = caption_p.add_run(f"{label}. {caption}")
    caption_run.bold = True
    caption_run.font.size = Pt(8.5)
    source = document.add_paragraph()
    source.add_run(f"Source: {path.relative_to(ROOT).as_posix()}").italic = True


def build_docx() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    document = Document()
    configure_document(document)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Supplementary Information").bold = True
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(
        "Residual-aware EEG learning for post-stroke proportional recovery"
    )

    document.add_heading("Supplementary Methods Notes", level=1)
    notes = [
        "All uncertainty estimates and statistical comparisons use subject-level LOSO predictions. Segment rows and seed rows are not treated as independent patients.",
        "Supplementary tables are compact manuscript views. Full source data and wide tables are provided in outputs/manuscript_package/ResidualAware_SSL_CNN_Source_Data.xlsx.",
        "Connectivity and topographic figures are model-explanation summaries and should be interpreted as hypothesis-generating.",
        "Participant-flow source notes summarize de-identified workbook counts and should not be interpreted as a complete adverse-event monitoring dataset.",
        "The PROBAST/TRIPOD+AI-oriented risk audit is conservative and highlights risks rather than treating current reporting completeness as evidence of low bias.",
        "Ethics approval, consent, EEG acquisition hardware, preprocessing filters, artifact rejection, recruitment dates, and final data-access route remain author-supplied fields.",
    ]
    for note in notes:
        document.add_paragraph(note, style="List Bullet")

    document.add_heading("Supplementary Figures", level=1)
    add_figure(document, *SUPPLEMENTARY_FIGURES[0], width=8.5)
    add_figure(document, *SUPPLEMENTARY_FIGURES[1], width=9.2)
    add_figure(document, *SUPPLEMENTARY_FIGURES[2], width=8.8)
    add_figure(document, *SUPPLEMENTARY_FIGURES[3], width=8.8)

    document.add_heading("Supplementary Tables", level=1)
    for label, caption, path, columns in SUPPLEMENTARY_TABLES:
        add_table(document, label, caption, path, columns)

    document.add_heading("Author-Supplied Fields Still Required", level=1)
    queries = [
        "Ethics approval institution, approval number, and consent/data-sharing wording.",
        "Recruitment dates, inclusion/exclusion criteria, stroke subtype, lesion-side definition, and timing from stroke to EEG/tACS.",
        "EEG acquisition hardware, electrode cap/montage description, reference, preprocessing filters, artifact rejection, and epoching criteria before feature extraction.",
        "Whether standardized conventional rehabilitation was delivered concurrently with the tACS protocol.",
        "Repository DOI and final access route for raw, processed, and derived data.",
    ]
    for query in queries:
        document.add_paragraph(query, style="List Number")

    document.save(OUTPUT_DOCX)
    print(OUTPUT_DOCX)


if __name__ == "__main__":
    build_docx()
