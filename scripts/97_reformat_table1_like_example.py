from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "output" / "doc" / "第二部分_材料与方法_学术论文风格重写稿_v2_公式编号修订版.docx"
OUTPUT_DOCX = ROOT / "output" / "doc" / "第二部分_材料与方法_学术论文风格重写稿_v2_公式编号修订版_表1重排版.docx"
BACKUP = ROOT / "output" / "doc" / "第二部分_材料与方法_学术论文风格重写稿_v2_表1修改前备份.docx"


BASELINE_TABLE = [
    ["指标", "全部患者", "比例恢复组", "恢复不良组", "P 值"],
    ["年龄", "64.03 ± 9.03", "64.00 ± 7.60", "66.11 ± 5.30", "0.49"],
    ["性别", "女性 15（51.7%）", "女性 6（60.0%）", "女性 6（66.7%）", "1.00"],
    ["病程", "5.86 ± 22.47（n=28）", "0.97 ± 0.52", "1.27 ± 0.71", "0.29"],
    ["卒中类型", "", "", "", ""],
    ["患侧/病灶侧", "左侧 14（48.3%）", "左侧 7（70.0%）", "左侧 4（44.4%）", "0.37"],
    ["治疗前 FMA-UE", "38.88 ± 23.72（n=25）", "59.60 ± 4.74", "19.22 ± 16.95", "<0.001"],
    ["治疗前 MBI", "60.00 ± 18.42（n=24）", "72.50 ± 13.59", "43.75 ± 15.29（n=8）", "<0.001"],
    ["基线 EEG 可用情况", "28（96.6%）", "10（100.0%）", "9（100.0%）", ""],
]


PAIRED_TABLE = [
    ["指标", "n", "治疗前", "治疗后", "改善量", "统计检验", "P 值"],
    ["FMA-UE", "19", "40.47 ± 23.83", "45.47 ± 23.23", "5.00 ± 4.07", "Wilcoxon signed-rank\n检验", "<0.001"],
    ["MBI", "19", "55.53 ± 19.92", "76.84 ± 21.49", "21.32 ± 14.22", "配对 t 检验", "<0.001"],
]


def set_run_font(run, size: float = 10.5, bold: bool = False) -> None:
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(size)
    run.bold = bold


def set_paragraph_text(paragraph, text: str, size: float = 12, bold: bool = False) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(8)


def insert_paragraph_before(element, text: str, size: float = 12, bold: bool = False):
    p = OxmlElement("w:p")
    element.addprevious(p)
    paragraph = element.getparent()._element_to_p(p) if hasattr(element.getparent(), "_element_to_p") else None
    # python-docx has no public wrapper for arbitrary inserted XML; wrap by
    # constructing through the document after insertion.
    return p


def make_paragraph_xml(text: str, size_half_points: int = 24, bold: bool = False, space_after: int = 160) -> OxmlElement:
    p = OxmlElement("w:p")
    p_pr = OxmlElement("w:pPr")
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), str(space_after))
    p_pr.append(spacing)
    p.append(p_pr)
    r = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "微软雅黑")
    fonts.set(qn("w:hAnsi"), "微软雅黑")
    fonts.set(qn("w:eastAsia"), "微软雅黑")
    r_pr.append(fonts)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(size_half_points))
    r_pr.append(sz)
    if bold:
        r_pr.append(OxmlElement("w:b"))
    r.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    r.append(t)
    p.append(r)
    return p


def clear_borders(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        tag = "w:" + edge
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        element.set(qn("w:val"), "nil")


def set_border(cell, edge: str, color: str = "E6E6E6", size: str = "4") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    element = tc_borders.find(qn("w:" + edge))
    if element is None:
        element = OxmlElement("w:" + edge)
        tc_borders.append(element)
    element.set(qn("w:val"), "single")
    element.set(qn("w:sz"), size)
    element.set(qn("w:space"), "0")
    element.set(qn("w:color"), color)


def set_cell_margins(cell, top: int = 120, start: int = 80, bottom: int = 120, end: int = 80) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_text(cell, text: str, bold: bool, align: WD_ALIGN_PARAGRAPH, size: float = 10.5) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    parts = text.split("\n")
    for idx, part in enumerate(parts):
        if idx:
            paragraph.add_run().add_break()
        run = paragraph.add_run(part)
        set_run_font(run, size=size, bold=bold)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def build_light_table(doc: Document, data: list[list[str]], widths: list[float], font_size: float) -> object:
    table = doc.add_table(rows=len(data), cols=len(data[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for r_idx, row in enumerate(data):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            clear_borders(cell)
            set_cell_margins(cell)
            set_cell_width(cell, widths[c_idx])
            align = WD_ALIGN_PARAGRAPH.LEFT if c_idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
            set_cell_text(cell, value, bold=(r_idx == 0), align=align, size=font_size)
            set_border(cell, "bottom", color="D9D9D9" if r_idx == 0 else "EEEEEE", size="5" if r_idx == 0 else "3")
    return table


def replace_table(old_table, new_table) -> None:
    old_tbl = old_table._tbl
    old_tbl.addprevious(new_table._tbl)
    old_tbl.getparent().remove(old_tbl)


def remove_paragraph(paragraph) -> None:
    element = paragraph._p
    element.getparent().remove(element)


def main() -> None:
    if not DOCX.exists():
        raise FileNotFoundError(DOCX)
    if not BACKUP.exists():
        shutil.copy2(DOCX, BACKUP)

    doc = Document(DOCX)
    if len(doc.tables) < 2:
        raise RuntimeError(f"Expected at least two tables, found {len(doc.tables)}")

    first_caption = None
    second_caption = None
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text.startswith("表 1A") or text.startswith("Table 1"):
            first_caption = paragraph
        elif text.startswith("表 1B"):
            second_caption = paragraph
    if first_caption is None or second_caption is None:
        raise RuntimeError("Could not find Table 1A/Table 1B captions")

    set_paragraph_text(first_caption, "Table 1 分成两个部分。", size=13, bold=True)

    first_table = doc.tables[0]
    second_table = doc.tables[1]

    # Insert the first explanatory line immediately before the first table.
    first_table._tbl.addprevious(
        make_paragraph_xml(
            "第一部分写“基线特征与分组比较”，只放治疗前就已经存在的变量，例如：",
            size_half_points=24,
            bold=False,
            space_after=120,
        )
    )

    new_first = build_light_table(doc, BASELINE_TABLE, [1.55, 1.55, 1.35, 1.35, 0.75], 10.5)
    replace_table(first_table, new_first)

    # Reuse the second caption paragraph as the explanatory sentence between the
    # two tables, matching the requested screenshot structure.
    set_paragraph_text(
        second_caption,
        "这里的目的才是说明：两组在年龄、性别、病程、患侧等基线变量上是否可比。",
        size=12,
        bold=False,
    )
    second_table._tbl.addprevious(
        make_paragraph_xml(
            "第二部分单独写“治疗前后临床量表变化”，放最终监督队列 19 例患者的配对结果：",
            size_half_points=24,
            bold=False,
            space_after=120,
        )
    )

    new_second = build_light_table(doc, PAIRED_TABLE, [0.8, 0.45, 1.1, 1.1, 1.05, 1.85, 0.65], 10.5)
    replace_table(second_table, new_second)

    try:
        doc.save(DOCX)
        saved = DOCX
    except PermissionError:
        doc.save(OUTPUT_DOCX)
        saved = OUTPUT_DOCX
    print(f"saved={saved}")
    print(f"backup={BACKUP}")


if __name__ == "__main__":
    main()
