from __future__ import annotations

import argparse
import json
from copy import deepcopy
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_TEMPLATE = ROOT / "docs" / "author_submission_metadata_template.json"
DEFAULT_WORKBOOK = ROOT / "outputs" / "manuscript_package" / "Author_Submission_Metadata_Intake.xlsx"
DEFAULT_OUTPUT_JSON = ROOT / "outputs" / "manuscript_package" / "author_submission_metadata_from_intake.json"
DEFAULT_REPORT = ROOT / "docs" / "author_metadata_intake_import_report.md"


def main() -> None:
    parser = argparse.ArgumentParser(description="Import completed author metadata intake workbook into JSON.")
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK, help="Completed XLSX intake workbook.")
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE, help="Base metadata JSON template.")
    parser.add_argument("--output-json", type=Path, default=DEFAULT_OUTPUT_JSON, help="Converted JSON output path.")
    parser.add_argument("--report", type=Path, default=DEFAULT_REPORT, help="Markdown import report path.")
    parser.add_argument(
        "--update-template",
        action="store_true",
        help="Overwrite the template JSON after import. Use only after reviewing the generated JSON.",
    )
    args = parser.parse_args()

    template_path = args.template.resolve()
    workbook_path = args.workbook.resolve()
    output_json_path = args.output_json.resolve()
    report_path = args.report.resolve()

    template = load_template(template_path)
    imported_values, field_evidence, warnings = read_intake_workbook(workbook_path)
    converted = merge_values(template, imported_values, field_evidence)
    stats = build_stats(template, converted, imported_values, field_evidence, warnings)

    output_json_path.parent.mkdir(parents=True, exist_ok=True)
    output_json_path.write_text(json.dumps(converted, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    report_path.write_text(build_report(workbook_path, template_path, output_json_path, stats), encoding="utf-8")

    if args.update_template:
        template_path.write_text(json.dumps(converted, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    print(report_path)
    print(output_json_path)
    print(f"imported_nonblank_values={stats['imported_nonblank_values']}")
    print(f"fields_with_imported_values={stats['fields_with_imported_values']}")
    print(f"fields_with_evidence_source={stats['fields_with_evidence_source']}")
    print(f"warnings={len(warnings)}")


def load_template(path: Path) -> dict[str, Any]:
    if not path.exists():
        raise FileNotFoundError(path)
    payload = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(payload, dict) or not isinstance(payload.get("fields"), dict):
        raise ValueError("Template must contain a top-level 'fields' object.")
    return payload


def read_intake_workbook(path: Path) -> tuple[dict[str, dict[str, str]], dict[str, str], list[str]]:
    if not path.exists():
        raise FileNotFoundError(path)
    workbook = load_workbook(path, read_only=False, data_only=True)
    if "Author_Input" not in workbook.sheetnames:
        raise ValueError("Workbook must contain an 'Author_Input' sheet.")
    imported_values: dict[str, dict[str, str]] = {}
    field_evidence: dict[str, str] = {}
    warnings: list[str] = []

    if "Minimal_Completion" in workbook.sheetnames:
        read_response_sheet(
            workbook["Minimal_Completion"],
            "Minimal_Completion",
            imported_values,
            field_evidence,
            warnings,
            overwrite=True,
        )
    read_response_sheet(
        workbook["Author_Input"],
        "Author_Input",
        imported_values,
        field_evidence,
        warnings,
        overwrite=False,
    )

    return imported_values, field_evidence, warnings


def read_response_sheet(
    sheet,
    sheet_name: str,
    imported_values: dict[str, dict[str, str]],
    field_evidence: dict[str, str],
    warnings: list[str],
    overwrite: bool,
) -> None:
    headers = [normalize_header(cell.value) for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    required_headers = {"field_id", "metadata_key", "author_response", "evidence_source"}
    missing_headers = sorted(required_headers - set(headers))
    if missing_headers:
        raise ValueError(f"{sheet_name} missing columns: {', '.join(missing_headers)}")

    index = {header: headers.index(header) for header in headers if header}

    for row_number, cells in enumerate(sheet.iter_rows(min_row=2), start=2):
        field_id = clean(cells[index["field_id"]].value)
        metadata_key = clean(cells[index["metadata_key"]].value)
        author_response = clean(cells[index["author_response"]].value)
        evidence_source = clean(cells[index["evidence_source"]].value)
        if not field_id and not metadata_key:
            continue
        if not field_id or not metadata_key:
            warnings.append(f"{sheet_name} row {row_number}: missing field_id or metadata_key.")
            continue
        if author_response:
            field_values = imported_values.setdefault(field_id, {})
            if overwrite or metadata_key not in field_values:
                field_values[metadata_key] = author_response
        if evidence_source:
            if field_id in field_evidence and field_evidence[field_id] != evidence_source:
                warnings.append(
                    f"{sheet_name} row {row_number}: multiple evidence_source values for {field_id}; using the first one."
                )
            if overwrite or field_id not in field_evidence:
                field_evidence.setdefault(field_id, evidence_source)


def merge_values(
    template: dict[str, Any],
    imported_values: dict[str, dict[str, str]],
    field_evidence: dict[str, str],
) -> dict[str, Any]:
    converted = deepcopy(template)
    fields = converted["fields"]
    for field_id, key_values in imported_values.items():
        if field_id not in fields or not isinstance(fields[field_id], dict):
            continue
        for key, value in key_values.items():
            if key in fields[field_id]:
                fields[field_id][key] = value
    for field_id, evidence in field_evidence.items():
        if field_id in fields and isinstance(fields[field_id], dict) and "evidence_source" in fields[field_id]:
            if not fields[field_id].get("evidence_source"):
                fields[field_id]["evidence_source"] = evidence
    return converted


def build_stats(
    template: dict[str, Any],
    converted: dict[str, Any],
    imported_values: dict[str, dict[str, str]],
    field_evidence: dict[str, str],
    warnings: list[str],
) -> dict[str, Any]:
    imported_count = sum(len(values) for values in imported_values.values())
    converted_fields = converted.get("fields", {})
    nonblank_by_field = {
        field_id: [key for key, value in values.items() if clean(value)]
        for field_id, values in converted_fields.items()
        if isinstance(values, dict)
    }
    return {
        "template_fields": len(template.get("fields", {})),
        "converted_fields": len(converted_fields),
        "imported_nonblank_values": imported_count,
        "fields_with_imported_values": sum(1 for values in imported_values.values() if values),
        "fields_with_evidence_source": len(field_evidence),
        "converted_nonblank_fields": sum(1 for keys in nonblank_by_field.values() if keys),
        "warnings": warnings,
    }


def build_report(workbook_path: Path, template_path: Path, output_json_path: Path, stats: dict[str, Any]) -> str:
    lines = [
        "# Author Metadata Intake Import Report",
        "",
        f"Workbook: `{workbook_path}`",
        f"Template JSON: `{template_path}`",
        f"Converted JSON: `{output_json_path}`",
        "",
        "This report summarizes import of author responses from the XLSX intake workbook into a JSON metadata file. If present, `Minimal_Completion` is read first and `Author_Input` is used as a fallback for still-blank keys. The import step does not verify the truth of supplied values. Run strict metadata validation before final manuscript insertion.",
        "",
        "## Summary",
        "",
        f"- Template fields: {stats['template_fields']}",
        f"- Converted fields: {stats['converted_fields']}",
        f"- Imported nonblank metadata values: {stats['imported_nonblank_values']}",
        f"- Fields with imported values: {stats['fields_with_imported_values']}",
        f"- Fields with evidence source: {stats['fields_with_evidence_source']}",
        f"- Converted fields with any nonblank value: {stats['converted_nonblank_fields']}",
        f"- Import warnings: {len(stats['warnings'])}",
        "",
        "## Next Commands",
        "",
        "```powershell",
        f"python scripts\\63_validate_author_submission_metadata.py --metadata \"{output_json_path}\" --strict",
        f"python scripts\\64_build_author_metadata_insertion_protocol.py --metadata \"{output_json_path}\" --strict",
        "```",
        "",
        "Use `--update-template` only after reviewing the converted JSON and confirming that the values are author-approved.",
        "",
        "## Warnings",
        "",
    ]
    if stats["warnings"]:
        lines.extend(f"- {warning}" for warning in stats["warnings"])
    else:
        lines.append("None.")
    lines.append("")
    return "\n".join(lines)


def normalize_header(value: Any) -> str:
    return clean(value).strip().lower()


def clean(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


if __name__ == "__main__":
    main()
