from __future__ import annotations

import sys
import csv
from collections import defaultdict
from pathlib import Path
from statistics import mean
from typing import Any

import openpyxl
import pandas as pd

PROJECT_ROOT = Path(__file__).resolve().parents[1]
SRC_ROOT = PROJECT_ROOT / "src"
if str(SRC_ROOT) not in sys.path:
    sys.path.insert(0, str(SRC_ROOT))

from eeg_recovery.config import load_path_config
from eeg_recovery.io.eeglab import read_eeglab_set_metadata
from eeg_recovery.io.index import build_eeg_file_index
from eeg_recovery.metadata.labels import load_supervised_label_table


CONFIG_PATH = PROJECT_ROOT / "configs" / "paths.example.yaml"
TABLE_DIR = PROJECT_ROOT / "results" / "tables"
DOC_DIR = PROJECT_ROOT / "docs"

EEG_RECORD_CSV = TABLE_DIR / "eeg_recording_metadata_audit.csv"
EEG_SUMMARY_CSV = TABLE_DIR / "eeg_recording_summary.csv"
WORKBOOK_AUDIT_CSV = TABLE_DIR / "clinical_workbook_structure_audit.csv"
EEG_AUDIT_MD = DOC_DIR / "eeg_metadata_audit.md"
AUTHOR_QUERY_MD = DOC_DIR / "author_information_request_table.md"


KEYWORD_GROUPS = {
    "date_or_timing": ("日期", "时间", "发病", "入院", "出院", "评估", "随访", "date", "time"),
    "stroke_or_diagnosis": ("卒中", "脑梗", "脑出血", "诊断", "病灶", "lesion", "stroke"),
    "ethics_or_consent": ("伦理", "审批", "知情", "同意", "ethic", "irb", "consent"),
    "eeg_or_protocol": ("脑电", "EEG", "采样", "电极", "滤波", "伪迹", "tACS", "刺激"),
}


def main() -> None:
    TABLE_DIR.mkdir(parents=True, exist_ok=True)
    DOC_DIR.mkdir(parents=True, exist_ok=True)

    config = load_path_config(CONFIG_PATH)
    labels = load_supervised_label_table(config)
    records = build_eeg_file_index(
        config.patient_eeg_root,
        config.health_eeg_root,
        labels["subject_id"].tolist(),
        validate_supervised_baseline=True,
    )

    eeg_rows = audit_eeg_records(records)
    write_csv(EEG_RECORD_CSV, eeg_rows)
    summary_rows = summarize_eeg_records(eeg_rows)
    write_csv(EEG_SUMMARY_CSV, summary_rows)

    workbook_rows = []
    for workbook_label, path in [
        ("integrity_workbook", config.patient_info_integrity_xlsx),
        ("clinical_workbook", config.patient_info_clinical_xlsx),
    ]:
        workbook_rows.extend(audit_workbook(workbook_label, path))
    write_csv(WORKBOOK_AUDIT_CSV, workbook_rows)

    write_eeg_audit_markdown(eeg_rows, summary_rows, workbook_rows)
    write_author_query_markdown(workbook_rows)
    print(f"Wrote {EEG_RECORD_CSV}")
    print(f"Wrote {EEG_SUMMARY_CSV}")
    print(f"Wrote {WORKBOOK_AUDIT_CSV}")
    print(f"Wrote {EEG_AUDIT_MD}")
    print(f"Wrote {AUTHOR_QUERY_MD}")


def audit_eeg_records(records: list[Any]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for record in records:
        metadata = read_eeglab_set_metadata(record.set_path)
        duration_sec = metadata.pnts * metadata.trials / metadata.srate
        rows.append(
            {
                "group": record.group,
                "subject_id": record.subject_id,
                "stage": record.stage,
                "state": record.state,
                "is_supervised_subject": int(record.is_supervised_subject),
                "nbchan": metadata.nbchan,
                "srate_hz": metadata.srate,
                "trials": metadata.trials,
                "points_per_trial": metadata.pnts,
                "duration_sec": round(duration_sec, 3),
                "xmin": metadata.xmin,
                "xmax": metadata.xmax,
                "channel_count": len(metadata.ch_names),
                "first_channel": metadata.ch_names[0] if metadata.ch_names else "",
                "last_channel": metadata.ch_names[-1] if metadata.ch_names else "",
                "resolved_fdt_exists": int(metadata.fdt_path.exists()),
            }
        )
    return rows


def summarize_eeg_records(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    grouped: dict[tuple[str, str, str], list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        grouped[(row["group"], row["stage"], row["state"])].append(row)

    summaries: list[dict[str, Any]] = []
    for (group, stage, state), items in sorted(grouped.items()):
        durations = [float(item["duration_sec"]) for item in items]
        pnts = [int(item["points_per_trial"]) for item in items]
        summaries.append(
            {
                "group": group,
                "stage": stage,
                "state": state,
                "n_records": len(items),
                "n_subjects": len({item["subject_id"] for item in items}),
                "srate_values_hz": join_unique(item["srate_hz"] for item in items),
                "nbchan_values": join_unique(item["nbchan"] for item in items),
                "trials_values": join_unique(item["trials"] for item in items),
                "duration_sec_mean": round(mean(durations), 3),
                "duration_sec_min": round(min(durations), 3),
                "duration_sec_max": round(max(durations), 3),
                "points_min": min(pnts),
                "points_max": max(pnts),
            }
        )
    return summaries


def audit_workbook(workbook_label: str, path: Path) -> list[dict[str, Any]]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows: list[dict[str, Any]] = []
    for sheet in workbook.worksheets:
        max_row = sheet.max_row or 0
        max_col = sheet.max_column or 0
        first_non_empty = find_first_non_empty_row(sheet)
        candidate_headers = candidate_header_cells(sheet)
        keyword_hits = keyword_counts(sheet, max_scan_rows=min(max_row, 80), max_scan_cols=min(max_col, 40))
        rows.append(
            {
                "workbook": workbook_label,
                "sheet_name": sheet.title,
                "max_row": max_row,
                "max_column": max_col,
                "first_non_empty_row": first_non_empty,
                "candidate_header_cells": " | ".join(candidate_headers[:30]),
                **keyword_hits,
            }
        )
    workbook.close()
    return rows


def find_first_non_empty_row(sheet: Any) -> int | str:
    for row_index, row in enumerate(
        sheet.iter_rows(min_row=1, max_row=min(sheet.max_row or 0, 50), values_only=True),
        start=1,
    ):
        if any(value not in (None, "") for value in row):
            return row_index
    return ""


def candidate_header_cells(sheet: Any) -> list[str]:
    header_tokens = (
        "患者ID",
        "编号",
        "年龄",
        "病程",
        "性别",
        "患病侧",
        "治疗前",
        "治疗后",
        "FMA",
        "MBI",
        "基线完整性",
        "睁眼",
        "闭眼",
    )
    fallback: list[str] = []
    for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row or 0, 20), values_only=True):
        values = [str(value).strip() for value in row[: min(len(row), 30)] if value not in (None, "")]
        if not values:
            continue
        if not fallback:
            fallback = [value for value in values if len(value) <= 40]
        joined = " ".join(values)
        if any(token in joined for token in header_tokens):
            return dedupe([value for value in values if len(value) <= 40])
    return dedupe(fallback)


def keyword_counts(sheet: Any, *, max_scan_rows: int, max_scan_cols: int) -> dict[str, int]:
    counters = {name: 0 for name in KEYWORD_GROUPS}
    for row in sheet.iter_rows(min_row=1, max_row=max_scan_rows, max_col=max_scan_cols, values_only=True):
        for value in row:
            if value is None:
                continue
            text = str(value)
            for name, keywords in KEYWORD_GROUPS.items():
                if any(keyword.lower() in text.lower() for keyword in keywords):
                    counters[name] += 1
    return counters


def write_eeg_audit_markdown(
    eeg_rows: list[dict[str, Any]],
    summary_rows: list[dict[str, Any]],
    workbook_rows: list[dict[str, Any]],
) -> None:
    supervised_baseline = [
        row
        for row in eeg_rows
        if row["group"] == "patient"
        and row["stage"] == "基线"
        and int(row["is_supervised_subject"]) == 1
    ]
    durations = [float(row["duration_sec"]) for row in supervised_baseline]
    pnts = [int(row["points_per_trial"]) for row in supervised_baseline]
    srate_values = join_unique(row["srate_hz"] for row in supervised_baseline)
    nbchan_values = join_unique(row["nbchan"] for row in supervised_baseline)
    trials_values = join_unique(row["trials"] for row in supervised_baseline)

    lines = [
        "# EEG And Clinical Source Metadata Audit",
        "",
        "This audit was generated from the current external data paths configured in `configs/paths.example.yaml`. It records non-identifying metadata only.",
        "",
        "## Supervised Baseline EEG Summary",
        "",
        f"- Records audited: {len(supervised_baseline)} baseline EO/EC records from supervised patients.",
        f"- Sampling-rate values: {srate_values} Hz.",
        f"- Channel-count values: {nbchan_values}.",
        f"- Trial-count values: {trials_values}.",
        f"- Recording duration: mean {round(mean(durations), 2)} s, range {round(min(durations), 2)}-{round(max(durations), 2)} s.",
        f"- Sample points per file: range {min(pnts)}-{max(pnts)}.",
        "",
        "## EEG Metadata Tables",
        "",
        "- Record-level audit: `results/tables/eeg_recording_metadata_audit.csv`.",
        "- Grouped summary: `results/tables/eeg_recording_summary.csv`.",
        "",
        "## Clinical Workbook Structure",
        "",
        "| Workbook | Sheet | Rows | Columns | Header candidates | Date/timing hits | Stroke/diagnosis hits | Ethics/consent hits | EEG/protocol hits |",
        "|---|---|---:|---:|---|---:|---:|---:|---:|",
    ]
    for row in workbook_rows:
        lines.append(
            "| {workbook} | {sheet_name} | {max_row} | {max_column} | {candidate_header_cells} | {date_or_timing} | {stroke_or_diagnosis} | {ethics_or_consent} | {eeg_or_protocol} |".format(
                **{key: markdown_cell(value) for key, value in row.items()}
            )
        )

    lines.extend(
        [
            "",
            "## Interpretation",
            "",
            "- The EEG metadata support reporting 250 Hz sampling, 62 retained channels, continuous single-trial files, and approximate baseline recording duration.",
            "- The workbook structure audit found whether candidate timing, diagnosis, ethics/consent, or protocol terms appear in the source sheets, but it does not expose identifiable patient-level values.",
            "- Ethics approval number, consent wording, acquisition hardware, reference scheme, filtering, artifact rejection, and exact recruitment criteria remain author-supplied items unless confirmed in a separate protocol document.",
            "",
        ]
    )
    EEG_AUDIT_MD.write_text("\n".join(lines), encoding="utf-8")


def write_author_query_markdown(workbook_rows: list[dict[str, Any]]) -> None:
    ethics_hits = sum(int(row["ethics_or_consent"]) for row in workbook_rows)
    timing_hits = sum(int(row["date_or_timing"]) for row in workbook_rows)
    diagnosis_hits = sum(int(row["stroke_or_diagnosis"]) for row in workbook_rows)
    protocol_hits = sum(int(row["eeg_or_protocol"]) for row in workbook_rows)

    rows = [
        ("Ethics approval institution and approval number", "Not verified in current audit" if ethics_hits == 0 else "Possible workbook terms found; author must verify exact wording"),
        ("Informed consent and data-sharing permission", "Not verified in current audit" if ethics_hits == 0 else "Possible workbook terms found; author must verify exact wording"),
        ("Recruitment dates and assessment timing", "Possible workbook terms found; author should extract final dates" if timing_hits else "Not verified in current audit"),
        ("Stroke subtype, lesion location, inclusion/exclusion criteria", "Possible workbook terms found; author should verify" if diagnosis_hits else "Not verified in current audit"),
        ("EEG acquisition hardware, electrode cap, reference, recording duration", "Duration estimated from `.set` metadata; hardware/cap/reference still required"),
        ("Preprocessing filters, artifact rejection, bad-channel handling", "Not verified in current audit"),
        ("Concurrent rehabilitation during tACS", "Possible protocol terms found; author should verify" if protocol_hits else "Not verified in current audit"),
        ("Repository DOI and data access route", "Not yet available"),
    ]

    lines = [
        "# Author Information Request Table",
        "",
        "Use this table to fill the remaining submission-critical fields. Do not delete uncertainty notes from the manuscript until the corresponding field is confirmed.",
        "",
        "| Required field | Current audit status |",
        "|---|---|",
    ]
    for field, status in rows:
        lines.append(f"| {field} | {status} |")
    lines.append("")
    AUTHOR_QUERY_MD.write_text("\n".join(lines), encoding="utf-8")


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    if not rows:
        return
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)


def join_unique(values: Any) -> str:
    return ";".join(str(value) for value in sorted(set(values), key=lambda item: str(item)))


def dedupe(values: list[str]) -> list[str]:
    seen = set()
    output = []
    for value in values:
        if value in seen:
            continue
        seen.add(value)
        output.append(value)
    return output


def markdown_cell(value: Any) -> str:
    return str(value).replace("|", "/").replace("\n", " ")


if __name__ == "__main__":
    main()
