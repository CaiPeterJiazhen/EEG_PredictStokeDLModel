from __future__ import annotations

import csv
import re
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
CLINICAL_WORKBOOK = Path(r"F:/CJZFile/EEG_M1/脑卒中患者信息记录表.xlsx")
INTEGRITY_WORKBOOK = Path(r"F:/CJZFile/EEG_M1/19例患者脑电数据完整性检查.xlsx")
PATIENT_EEG_ROOT = Path(r"F:/CJZFile/EEG_M1/Patient_tACS_M1_RestingStateEEG_afterProcess")
OUTPUT_CSV = ROOT / "results" / "tables" / "participant_flow_safety_source_notes.csv"
OUTPUT_MD = ROOT / "docs" / "participant_flow_and_safety_source_notes.md"


def main() -> None:
    m1_records = load_m1_records(CLINICAL_WORKBOOK)
    supervised_ids = load_supervised_ids(INTEGRITY_WORKBOOK)
    eeg_ids = load_eeg_indexed_ids(PATIENT_EEG_ROOT)

    m1_ids = {record["subject_id"] for record in m1_records}
    eeg_m1_ids = sorted(eeg_ids & m1_ids)
    no_current_eeg_ids = sorted(m1_ids - eeg_ids)
    supervised_eeg_ids = sorted(supervised_ids & eeg_ids)
    nonsupervised_eeg_ids = sorted(eeg_ids - supervised_ids)

    by_id = {record["subject_id"]: record for record in m1_records}
    nonsupervised_records = [by_id[subject_id] for subject_id in nonsupervised_eeg_ids if subject_id in by_id]
    no_eeg_records = [by_id[subject_id] for subject_id in no_current_eeg_ids if subject_id in by_id]
    note_counts = summarize_source_notes(nonsupervised_records + no_eeg_records)

    rows = build_rows(
        n_m1=len(m1_records),
        n_eeg=len(eeg_m1_ids),
        n_no_current_eeg=len(no_current_eeg_ids),
        n_supervised=len(supervised_eeg_ids),
        n_nonsupervised=len(nonsupervised_eeg_ids),
        n_label1=10,
        n_label0=9,
        nonsupervised_records=nonsupervised_records,
        no_eeg_records=no_eeg_records,
        note_counts=note_counts,
    )

    OUTPUT_CSV.parent.mkdir(parents=True, exist_ok=True)
    with OUTPUT_CSV.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=[
                "row_type",
                "category",
                "n",
                "denominator",
                "source_basis",
                "manuscript_use",
                "notes",
            ],
        )
        writer.writeheader()
        writer.writerows(rows)

    OUTPUT_MD.write_text(build_markdown(rows), encoding="utf-8")
    print(OUTPUT_CSV)
    print(OUTPUT_MD)
    print(f"participant_flow_rows={len(rows)}")


def load_m1_records(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet = workbook.active
    header_row = None
    start_row = None
    end_row = sheet.max_row + 1
    for row_index in range(1, sheet.max_row + 1):
        first = normalize_cell(sheet.cell(row_index, 1).value)
        if first == "M1组":
            header_row = row_index + 1
            start_row = row_index + 2
            continue
        if first == "健康受试者组" and start_row is not None:
            end_row = row_index
            break
    if header_row is None or start_row is None:
        raise ValueError(f"Could not locate M1 patient block in {path}")

    headers = [normalize_cell(sheet.cell(header_row, column).value) for column in range(1, sheet.max_column + 1)]
    records: list[dict[str, Any]] = []
    for row_index in range(start_row, end_row):
        row_values = [sheet.cell(row_index, column).value for column in range(1, sheet.max_column + 1)]
        if not any(value is not None and str(value).strip() for value in row_values):
            continue
        record = dict(zip(headers, row_values))
        subject_id = normalize_subject_id(record.get("编号"))
        if not subject_id:
            continue
        records.append(
            {
                "subject_id": subject_id,
                "fma_pre": record.get("治疗前FMA"),
                "fma_post": record.get("治疗后FMA"),
                "missing_data": normalize_cell(record.get("缺少数据")),
                "drop_reason": normalize_cell(record.get("脱落原因")),
            }
        )
    return records


def load_supervised_ids(path: Path) -> set[str]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet = workbook.active
    subject_ids: set[str] = set()
    for row_index in range(2, sheet.max_row + 1):
        subject_id = normalize_subject_id(sheet.cell(row_index, 1).value)
        if subject_id:
            subject_ids.add(subject_id)
    return subject_ids


def load_eeg_indexed_ids(root: Path) -> set[str]:
    subject_ids: set[str] = set()
    for path in root.rglob("*"):
        if path.is_dir():
            subject_id = normalize_subject_id(path.name)
            if subject_id:
                subject_ids.add(subject_id)
    return subject_ids


def summarize_source_notes(records: list[dict[str, Any]]) -> Counter[str]:
    counts: Counter[str] = Counter()
    for record in records:
        reason = record["drop_reason"]
        missing = record["missing_data"]
        fma_pre = record["fma_pre"]
        fma_post = record["fma_post"]
        if "未接受治疗" in reason:
            counts["Did not receive treatment"] += 1
        elif "出院" in reason:
            counts["Discharged before complete follow-up"] += 1
        elif "认知" in reason or "依从性差" in reason:
            counts["Poor compliance or cognitive/communication difficulty"] += 1
        elif "脑电帽" in reason or "闷热" in reason:
            counts["EEG-cap heat/discomfort note"] += 1
        elif "高血压" in reason:
            counts["Post-session discomfort with next-day hypertension note"] += 1
        elif "不舒服" in reason and "核磁" in reason:
            counts["MRI-related discomfort/no desire to enroll note"] += 1
        elif "不舒服" in reason:
            counts["Post-session discomfort note"] += 1
        elif missing == "无脑电":
            counts["No current EEG data in project directory"] += 1
        elif is_number(fma_pre) and is_number(fma_post) and float(fma_pre) >= 66:
            counts["Complete FMA with ceiling-level baseline score"] += 1
        elif missing or reason:
            counts["Other source-note reason"] += 1
        else:
            counts["Not in final supervised list; no source-note reason in workbook"] += 1
    return counts


def build_rows(
    *,
    n_m1: int,
    n_eeg: int,
    n_no_current_eeg: int,
    n_supervised: int,
    n_nonsupervised: int,
    n_label1: int,
    n_label0: int,
    nonsupervised_records: list[dict[str, Any]],
    no_eeg_records: list[dict[str, Any]],
    note_counts: Counter[str],
) -> list[dict[str, str | int]]:
    rows: list[dict[str, str | int]] = [
        flow_row("M1 patient records in clinical source workbook", n_m1, n_m1, "Cohort source frame"),
        flow_row("Current EEG-indexed M1 patient pool", n_eeg, n_m1, "Unlabeled/self-supervised patient EEG pool"),
        flow_row("Clinical workbook entries without current indexed EEG", n_no_current_eeg, n_m1, "Excluded from EEG analyses"),
        flow_row("Final labeled supervised cohort", n_supervised, n_eeg, "Patient-level LOSO model evaluation"),
        flow_row("EEG-indexed patients not used for supervised labels", n_nonsupervised, n_eeg, "Unlabeled/descriptive pool only"),
        flow_row("Proportional-recovery label", n_label1, n_supervised, "Outcome class in supervised cohort"),
        flow_row("Poor-recovery label", n_label0, n_supervised, "Outcome class in supervised cohort"),
    ]

    rows.append(
        source_note_row(
            "Non-supervised EEG-indexed entries with missing-data or discontinuation notes",
            sum(
                1
                for record in nonsupervised_records
                if record["missing_data"] or record["drop_reason"]
            ),
            len(nonsupervised_records),
            "Participant-flow interpretation only; not a complete adverse-event summary.",
        )
    )
    rows.append(
        source_note_row(
            "Clinical workbook entries without current EEG but with missing-data or discontinuation notes",
            sum(1 for record in no_eeg_records if record["missing_data"] or record["drop_reason"]),
            max(len(no_eeg_records), 1),
            "Participant-flow interpretation only; not a complete adverse-event summary.",
        )
    )
    for category, count in sorted(note_counts.items()):
        rows.append(
            source_note_row(
                category,
                count,
                n_nonsupervised + n_no_current_eeg,
                "Source workbook reason category; requires author review before formal safety reporting.",
            )
        )
    return rows


def flow_row(category: str, n: int, denominator: int, manuscript_use: str) -> dict[str, str | int]:
    return {
        "row_type": "participant_flow",
        "category": category,
        "n": n,
        "denominator": denominator,
        "source_basis": "M1 clinical source workbook, patient EEG directory index, and 19-patient integrity workbook",
        "manuscript_use": manuscript_use,
        "notes": "Counts are de-identified and derived from current project files.",
    }


def source_note_row(category: str, n: int, denominator: int, notes: str) -> dict[str, str | int]:
    return {
        "row_type": "source_note_category",
        "category": category,
        "n": n,
        "denominator": denominator,
        "source_basis": "M1 clinical source workbook missing-data/drop-reason notes",
        "manuscript_use": "Author-facing participant-flow and safety-source review",
        "notes": notes,
    }


def build_markdown(rows: list[dict[str, str | int]]) -> str:
    lines = [
        "# Participant Flow And Safety Source Notes",
        "",
        "This audit summarizes de-identified participant-flow counts and source-workbook note categories for manuscript Figure 1 and author safety review. It does not replace an investigator-adjudicated adverse-event table.",
        "",
        "| Row type | Category | n | Denominator | Manuscript use | Notes |",
        "|---|---|---:|---:|---|---|",
    ]
    for row in rows:
        lines.append(
            f"| {row['row_type']} | {escape(row['category'])} | {row['n']} | {row['denominator']} | {escape(row['manuscript_use'])} | {escape(row['notes'])} |"
        )
    lines.extend(
        [
            "",
            "## Interpretation",
            "",
            "- These counts support transparent cohort-flow reporting and the Figure 1 participant-flow panel.",
            "- The source workbook contains missing-data and discontinuation notes, including discomfort-related entries, but it does not document the adverse-event monitoring method, severity, causality, or complete denominator required for final safety reporting.",
            "- Formal safety/tolerability wording remains author-required and must be confirmed against clinical records, ethics documents, or the study protocol.",
            "",
        ]
    )
    return "\n".join(lines)


def normalize_cell(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def normalize_subject_id(value: Any) -> str:
    match = re.search(r"sub0*(\d+)", normalize_cell(value), flags=re.IGNORECASE)
    if not match:
        return ""
    return f"sub{int(match.group(1)):02d}"


def is_number(value: Any) -> bool:
    try:
        float(value)
    except (TypeError, ValueError):
        return False
    return True


def escape(value: Any) -> str:
    return str(value).replace("|", "\\|").replace("\n", " ")


if __name__ == "__main__":
    main()
