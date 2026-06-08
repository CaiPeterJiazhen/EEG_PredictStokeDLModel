from __future__ import annotations

import csv
import os
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
MANUSCRIPT_MD = Path(
    os.environ.get(
        "MANUSCRIPT_MD",
        ROOT / "docs" / "manuscript_residual_aware_ssl_cnn_nature_polished.md",
    )
)
OUTPUT_DIR = ROOT / "output" / "doc"
OUTPUT_DOCX = Path(
    os.environ.get(
        "OUTPUT_DOCX",
        OUTPUT_DIR / "ResidualAware_SSL_CNN_Nature_Manuscript.docx",
    )
)

FIGURES = {
    "Figure 1": ROOT / "results" / "figures" / "nature" / "figure1_study_design_model.png",
    "Figure 2": ROOT / "results" / "figures" / "nature" / "figure2_performance_calibration.png",
    "Figure 3": ROOT / "results" / "figures" / "nature" / "figure3_robustness_ablation.png",
    "Figure 4": ROOT / "results" / "figures" / "nature" / "figure4_explainability_neurophysiology.png",
    "Supplementary Figure 1": ROOT / "results" / "figures" / "nature" / "supplementary_error_subjects.png",
    "Supplementary Figure 2": ROOT
    / "results"
    / "figures"
    / "explainability"
    / "mne_wpli_connectivity"
    / "mne_wpli_connectivity_contact_sheet.png",
}

TABLES = {
    "Table 1": ROOT / "results" / "tables" / "patient_characteristics_table.csv",
    "Table 2": ROOT / "results" / "tables" / "model_performance_main_table.csv",
    "Table 3": ROOT / "results" / "tables" / "table3_ablation.csv",
    "Table 4": ROOT / "results" / "tables" / "explainability_key_findings_table.csv",
}


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(7)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def format_number(value: str) -> str:
    if value is None:
        return ""
    value = str(value)
    if value == "" or re.search(r"[A-Za-z%/\-:]", value):
        return value
    try:
        number = float(value)
    except ValueError:
        return value
    if abs(number) >= 10:
        return f"{number:.2f}"
    return f"{number:.3f}"


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def compact_table_rows(label: str, rows: list[dict[str, str]]) -> tuple[list[str], list[list[str]]]:
    if label == "Table 1":
        columns = ["variable", "all", "proportional_label1", "poor_recovery_label0"]
        return columns, [[format_number(row.get(col, "")) for col in columns] for row in rows]

    if label == "Table 2":
        keep = {
            "ML_EEG_updated_no_selector_logistic_l1",
            "no_SSL_CNN_updated_sub05_sub28_seedensemble10",
            "residual_aware_SSL_CNN_seedmean10",
        }
        columns = ["model", "accuracy", "balanced_accuracy", "roc_auc", "pr_auc", "brier_score"]
        filtered = [row for row in rows if row.get("source_model") in keep]
        return columns, [[format_number(row.get(col, "")) for col in columns] for row in filtered]

    if label == "Table 3":
        columns = ["ablation_block", "model_key", "row_type", "accuracy", "roc_auc", "pr_auc", "brier_score"]
        selected = []
        for row in rows:
            if row.get("ablation_block") == "core" and row.get("row_type") in {"reported", "seedmean10", "ensemble10"}:
                selected.append(row)
            elif row.get("ablation_block") == "modality_state_band" and row.get("ablation_name") in {
                "psd_only",
                "wpli_only",
                "ec_only",
                "eo_only",
                "motor_wpli_edges",
            }:
                row = dict(row)
                row["model_key"] = row.get("ablation_name", "")
                selected.append(row)
        return columns, [[format_number(row.get(col, "")) for col in columns] for row in selected]

    if label == "Table 4":
        columns = ["finding_type", "finding", "metric", "value", "stability_or_validation"]
        return columns, [[format_number(row.get(col, "")) for col in columns] for row in rows[:12]]

    columns = list(rows[0].keys()) if rows else []
    return columns, [[format_number(row.get(col, "")) for col in columns] for row in rows]


def add_data_table(document: Document, label: str, caption: str) -> None:
    rows = read_csv(TABLES[label])
    columns, body = compact_table_rows(label, rows)
    paragraph = document.add_paragraph()
    run = paragraph.add_run(f"{label}. {caption}")
    run.bold = True
    run.font.size = Pt(9)

    table = document.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for index, col in enumerate(columns):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, "D9E2F3")
        set_cell_text(cell, col, bold=True)

    for row_values in body:
        cells = table.add_row().cells
        for index, value in enumerate(row_values):
            set_cell_text(cells[index], value)

    document.add_paragraph()


def add_figure(document: Document, label: str, caption: str) -> None:
    path = FIGURES[label]
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(path), width=Inches(6.3))
    caption_p = document.add_paragraph()
    caption_run = caption_p.add_run(f"{label}. {caption}")
    caption_run.bold = True
    caption_run.font.size = Pt(9)
    document.add_paragraph()


def add_inline_runs(paragraph, text: str) -> None:
    # Keep markdown-style inline code readable in Word without a full Markdown parser.
    parts = re.split(r"(`[^`]+`)", text)
    for part in parts:
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Courier New"
            run.font.size = Pt(9)
        else:
            paragraph.add_run(part)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.05

    for style_name, size in [("Heading 1", 15), ("Heading 2", 12), ("Heading 3", 10.5), ("Title", 18)]:
        style = styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True


def build_docx() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    text = MANUSCRIPT_MD.read_text(encoding="utf-8")
    lines = text.splitlines()
    caption_lookup: dict[str, str] = {}
    for raw_line in lines:
        line = raw_line.strip()
        for label in [
            "Figure 1",
            "Figure 2",
            "Figure 3",
            "Figure 4",
            "Supplementary Figure 1",
            "Supplementary Figure 2",
        ]:
            if line.startswith(f"{label}."):
                caption_lookup[label] = line.replace(f"{label}. ", "")

    document = Document()
    configure_document(document)

    inserted: set[str] = set()
    current_section = ""
    skip_source_lines = False

    for raw_line in lines:
        line = raw_line.strip()
        if not line:
            continue

        if line == "## Tables":
            skip_source_lines = True
            current_section = "Tables"
            continue
        if line == "## Figure legends":
            skip_source_lines = False
            current_section = "Figure legends"
        elif line.startswith("## "):
            skip_source_lines = False
            current_section = line[3:]
        elif line.startswith("### "):
            current_section = line[4:]

        if skip_source_lines:
            continue

        if line.startswith("# "):
            title = document.add_paragraph(style="Title")
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            title.add_run(line[2:]).bold = True
            document.add_paragraph("Draft manuscript generated from the current project analysis outputs.", style="Normal")
            continue

        if line.startswith("## "):
            document.add_heading(line[3:], level=1)
            continue

        if line.startswith("### "):
            document.add_heading(line[4:], level=2)
            continue

        if line.startswith("Figure 1."):
            if "Figure 1" not in inserted:
                add_figure(document, "Figure 1", caption_lookup.get("Figure 1", line.replace("Figure 1. ", "")))
                inserted.add("Figure 1")
            continue
        if line.startswith("Figure 2."):
            if "Figure 2" not in inserted:
                add_figure(document, "Figure 2", caption_lookup.get("Figure 2", line.replace("Figure 2. ", "")))
                inserted.add("Figure 2")
            continue
        if line.startswith("Figure 3."):
            if "Figure 3" not in inserted:
                add_figure(document, "Figure 3", caption_lookup.get("Figure 3", line.replace("Figure 3. ", "")))
                inserted.add("Figure 3")
            continue
        if line.startswith("Figure 4."):
            if "Figure 4" not in inserted:
                add_figure(document, "Figure 4", caption_lookup.get("Figure 4", line.replace("Figure 4. ", "")))
                inserted.add("Figure 4")
            continue
        if line.startswith("Supplementary Figure 1."):
            if "Supplementary Figure 1" not in inserted:
                add_figure(
                    document,
                    "Supplementary Figure 1",
                    caption_lookup.get("Supplementary Figure 1", line.replace("Supplementary Figure 1. ", "")),
                )
                inserted.add("Supplementary Figure 1")
            continue
        if line.startswith("Supplementary Figure 2."):
            if "Supplementary Figure 2" not in inserted:
                add_figure(
                    document,
                    "Supplementary Figure 2",
                    caption_lookup.get("Supplementary Figure 2", line.replace("Supplementary Figure 2. ", "")),
                )
                inserted.add("Supplementary Figure 2")
            continue

        paragraph = document.add_paragraph()
        if re.match(r"^\d+\. ", line):
            paragraph.style = document.styles["List Number"]
            line = re.sub(r"^\d+\. ", "", line)
        add_inline_runs(paragraph, line)

        if current_section == "Introduction" and line.startswith("The objective of this study") and "Figure 1" not in inserted:
            add_figure(document, "Figure 1", caption_lookup.get("Figure 1", "Study design and model architecture."))
            inserted.add("Figure 1")

        if "Descriptive cohort statistics are provided in Table 1" in line and "Table 1" not in inserted:
            add_data_table(document, "Table 1", "Cohort characteristics for the supervised 19-patient cohort.")
            inserted.add("Table 1")
        if "Table 2; Figure 2" in line and "Figure 2" not in inserted:
            add_figure(document, "Figure 2", caption_lookup.get("Figure 2", "Primary performance, uncertainty, and calibration."))
            inserted.add("Figure 2")
        if "Table 2; Figure 2" in line and "Table 2" not in inserted:
            add_data_table(document, "Table 2", "Main patient-level model performance.")
            inserted.add("Table 2")
        if "Table 3; Figure 3" in line and "Figure 3" not in inserted:
            add_figure(document, "Figure 3", caption_lookup.get("Figure 3", "Robustness, ablation, and threshold sensitivity."))
            inserted.add("Figure 3")
        if "Table 3; Figure 3" in line and "Table 3" not in inserted:
            add_data_table(document, "Table 3", "Selected ablation and feature-family results.")
            inserted.add("Table 3")
        if current_section == "Model explanation and EEG biomarker localization" and "Figure 4" in line and "Figure 4" not in inserted:
            add_figure(document, "Figure 4", caption_lookup.get("Figure 4", "Explainability and neurophysiological interpretation."))
            inserted.add("Figure 4")
            add_data_table(document, "Table 4", "Selected explainability biomarkers and validation notes.")
            inserted.add("Table 4")

    document.add_section(WD_SECTION.NEW_PAGE)
    document.add_heading("Build Notes", level=1)
    notes = [
        f"This DOCX is generated from {MANUSCRIPT_MD.relative_to(ROOT).as_posix()}.",
        "Tables are compact manuscript views; full source tables remain in results/tables/.",
        "Figures are embedded as PNG drafting previews; TIFF/PDF/SVG exports are available in results/figures/nature/.",
        "Author-supplied ethics, intervention, and data-sharing details remain required before submission.",
    ]
    for note in notes:
        document.add_paragraph(note, style="List Bullet")

    document.save(OUTPUT_DOCX)
    print(OUTPUT_DOCX)


if __name__ == "__main__":
    build_docx()
