from __future__ import annotations

import json
import re
import csv
from collections import Counter
from copy import deepcopy
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE_JSON = ROOT / "docs" / "author_submission_metadata_template.json"
OUTPUT_JSON = ROOT / "outputs" / "manuscript_package" / "author_submission_metadata_project_prefill.json"
OUTPUT_MD = ROOT / "docs" / "author_metadata_project_prefill_report.md"
PATIENT_RAW_CNT_DIR = Path(r"F:\CJZFile\EEG_M1\Patient_tACS_M1_EEG")
PATIENT_CLINICAL_WORKBOOK = Path(r"F:\CJZFile\EEG_M1\脑卒中患者信息记录表.xlsx")
CLINICAL_RECORD_WORKBOOK = Path(r"F:\CJZFile\EEG_M1\M1组病历记录表.xlsx")
PATIENT_PDF_AUDIT_CSV = ROOT / "results" / "tables" / "patient_record_pdf_text_audit.csv"


PREFILLS: dict[str, dict[str, dict[str, str]]] = {
    "tacs_device_electrodes": {
        "confirmed_target_frequency_intensity_duration_sessions": {
            "value": (
                "Project records define a common tACS protocol: stimulation over contralateral M1, "
                "C3 for right-hand impairment and C4 for left-hand impairment, 20 Hz, 1000 microampere, "
                "20 min per session, once daily for 14 sessions over 2 weeks."
            ),
            "evidence": "docs/methods_detail_provenance.md; tacs_eeg_proportional_recovery_project_design.md",
        }
    },
    "study_site_dates_design": {
        "follow_up_or_last_assessment_window": {
            "value": (
                "The project design defines outcome assessment immediately after the final tACS session; "
                "no longer-term follow-up window is documented in the current project files."
            ),
            "evidence": "tacs_eeg_proportional_recovery_project_design.md; docs/methods_detail_provenance.md",
        }
    },
    "eligibility_stroke_timing": {
        "time_since_stroke_to_eeg": {
            "value": (
                "The M1-group clinical source workbook contains a disease-duration field. In the final "
                "19-patient supervised cohort, disease duration averaged 36.3 days, median 33 days, range 17-83 days. "
                "The source workbook does not state whether this interval is specifically stroke-to-EEG, "
                "stroke-to-tACS, or stroke-to-recruitment; author confirmation is required before final wording."
            ),
            "evidence": "F:/CJZFile/EEG_M1/脑卒中患者信息记录表.xlsx; F:/CJZFile/EEG_M1/19例患者脑电数据完整性检查.xlsx; docs/cohort_characteristics.md",
        }
    },
    "fma_assessors_timing": {
        "baseline_assessment_timing": {
            "value": "Baseline FMA-UE was assessed before the tACS intervention and before baseline EEG feature extraction.",
            "evidence": "docs/methods_detail_provenance.md; docs/project_context.md",
        },
        "post_treatment_assessment_timing": {
            "value": "Post-treatment FMA-UE was assessed after the final tACS session according to the current project context.",
            "evidence": "docs/methods_detail_provenance.md; docs/project_context.md",
        },
        "fma_ue_version_or_scoring_reference": {
            "value": (
                "The current label-generation code and project context use the FMA-UE upper-extremity maximum score of 66 "
                "for proportional-recovery calculations."
            ),
            "evidence": "docs/project_context.md; src/eeg_recovery/metadata/labels.py",
        },
    },
    "eeg_hardware_reference_impedance": {
        "original_channel_count": {
            "value": (
                "Project design describes 64-channel EEG; the current analysis retained 62 channels after M1/M2 removal."
            ),
            "evidence": "docs/methods_detail_provenance.md; docs/project_context.md; configs/channel_mapping.yaml",
        }
    },
    "resting_state_instructions": {
        "eyes_open_eyes_closed_order": {
            "value": (
                "Current project context maps baseline *1.set files to eyes-open resting state and *2.set files to "
                "eyes-closed resting state. The actual acquisition instructions and whether this file order exactly "
                "matches acquisition order require author confirmation."
            ),
            "evidence": "docs/project_context.md; docs/eeg_metadata_audit.md",
        },
        "target_duration_per_condition": {
            "value": (
                "Across the 38 supervised baseline EO/EC files, recording duration averaged 188.4 s "
                "and ranged from 101.0 to 247.8 s."
            ),
            "evidence": "results/tables/eeg_recording_metadata_audit.csv; docs/eeg_metadata_audit.md",
        }
    },
    "tacs_safety_adverse_events": {
        "withdrawals_or_discontinuations": {
            "value": (
                "The M1-group patient-information source workbook records 29 M1-group entries and 9 rows with missing "
                "EEG or follow-up data. Recorded reasons include non-treatment, discharge, poor compliance or cognitive/"
                "communication difficulty, EEG-cap heat discomfort, one post-session discomfort entry, one MRI-related "
                "discomfort entry, and one entry describing post-session discomfort followed by hypertension the next day. "
                "This is a participant-flow and source-note summary, not a complete adverse-event monitoring statement."
            ),
            "evidence": "F:/CJZFile/EEG_M1/脑卒中患者信息记录表.xlsx",
        }
    },
    "raw_eeg_preprocessing": {
        "eeglab_set_fdt_export_rules": {
            "value": (
                "Manuscript analyses start from project-provided preprocessed EEGLAB .set/.fdt files. "
                "After loading those files, the feature pipeline performs no additional temporal filtering, "
                "artifact rejection, channel interpolation, or bad-channel removal."
            ),
            "evidence": "docs/methods_detail_provenance.md; docs/methods_gap_resolution_from_project_files.md; src/eeg_recovery/io/eeglab.py",
        },
        "segmentation_or_continuous_export_rules": {
            "value": "The current indexed EEGLAB files are continuous one-trial recordings loaded as channels by samples.",
            "evidence": "docs/methods_detail_provenance.md; docs/eeg_metadata_audit.md",
        },
    },
    "data_repository_doi_scope": {
        "public_data_scope": {
            "value": (
                "Prepared public/derived materials include subject-level derived analysis tables, locked LOSO predictions, "
                "bootstrap/permutation/paired-comparison outputs, figure-source summaries, validation and artifact-quality "
                "audit tables, author-field replacement maps, reproducibility scripts, and source-data workbooks."
            ),
            "evidence": "docs/data_availability_and_fair_audit.md; docs/repository_readme_for_deposit.md; outputs/manuscript_package/ResidualAware_SSL_CNN_Source_Data.xlsx",
        },
        "restricted_data_scope": {
            "value": (
                "Raw EEG recordings, minimally processed EEG files, identifiable clinical source records, and directly "
                "linkable participant-level source files remain restricted pending ethics, consent, privacy, and data-use confirmation."
            ),
            "evidence": "docs/data_availability_and_fair_audit.md; docs/repository_readme_for_deposit.md",
        },
    },
    "code_repository_license": {
        "environment_or_runtime_notes": {
            "value": (
                "The local package includes scripts for feature extraction, model evaluation, statistical validation, figure generation, "
                "source-data assembly, manuscript generation, artifact audit, package assembly, and final submission gates."
            ),
            "evidence": "docs/repository_readme_for_deposit.md; outputs/submission_package_20260601/submission_package_manifest.csv",
        }
    },
}


AUTHOR_REQUIRED_FIELDS = [
    "ethics_approval",
    "informed_consent",
    "trial_or_study_registration",
    "study_site_dates_design",
    "eligibility_stroke_timing",
    "concurrent_rehabilitation",
    "tacs_safety_adverse_events",
    "target_journal_reference_style",
    "author_list_affiliations",
    "author_contributions",
    "funding_competing_acknowledgements",
]


def main() -> None:
    template = json.loads(TEMPLATE_JSON.read_text(encoding="utf-8"))
    prefilled = deepcopy(template)
    prefill_payload = merge_prefills(PREFILLS, build_dynamic_prefills())
    applied_rows: list[dict[str, str]] = []
    for field_id, key_map in prefill_payload.items():
        field = prefilled["fields"].setdefault(field_id, {})
        evidence_values: list[str] = []
        for key, payload in key_map.items():
            field[key] = payload["value"]
            evidence_values.append(payload["evidence"])
            applied_rows.append(
                {
                    "field_id": field_id,
                    "metadata_key": key,
                    "value": payload["value"],
                    "evidence": payload["evidence"],
                }
            )
        if evidence_values and not field.get("evidence_source"):
            field["evidence_source"] = "; ".join(sorted(set(evidence_values)))

    prefilled["_instructions"]["purpose"] = (
        "Project-evidence prefill draft. Author approval is still required before final manuscript replacement."
    )
    prefilled["_instructions"]["do_not_guess"] = (
        "Values in this file are limited to project-audited evidence. Blank fields remain author-required."
    )

    OUTPUT_JSON.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT_JSON.write_text(json.dumps(prefilled, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    OUTPUT_MD.write_text(build_report(applied_rows), encoding="utf-8")
    print(OUTPUT_JSON)
    print(OUTPUT_MD)
    print(f"prefilled_values={len(applied_rows)}")
    print(f"fields_with_prefill={len(set(row['field_id'] for row in applied_rows))}")
    print(f"author_required_fields={len(AUTHOR_REQUIRED_FIELDS)}")


def merge_prefills(*sources: dict[str, dict[str, dict[str, str]]]) -> dict[str, dict[str, dict[str, str]]]:
    merged: dict[str, dict[str, dict[str, str]]] = {}
    for source in sources:
        for field_id, key_map in source.items():
            merged.setdefault(field_id, {}).update(key_map)
    return merged


def build_dynamic_prefills() -> dict[str, dict[str, dict[str, str]]]:
    dynamic: dict[str, dict[str, dict[str, str]]] = {}
    cnt_window = collect_patient_cnt_window()
    if cnt_window:
        dynamic.setdefault("study_site_dates_design", {})["patient_raw_cnt_recording_window_project_evidence"] = {
            "value": cnt_window,
            "evidence": "F:/CJZFile/EEG_M1/Patient_tACS_M1_EEG raw CNT headers and file inventory",
        }
    clinical_summary = collect_clinical_workbook_summary()
    if clinical_summary:
        dynamic.setdefault("eligibility_stroke_timing", {})["source_workbook_clinical_fields_project_evidence"] = {
            "value": clinical_summary["clinical_fields"],
            "evidence": "F:/CJZFile/EEG_M1/脑卒中患者信息记录表.xlsx; F:/CJZFile/EEG_M1/M1组病历记录表.xlsx",
        }
        dynamic.setdefault("fma_assessors_timing", {})["source_workbook_scale_fields_project_evidence"] = {
            "value": clinical_summary["scale_fields"],
            "evidence": "F:/CJZFile/EEG_M1/脑卒中患者信息记录表.xlsx; F:/CJZFile/EEG_M1/M1组病历记录表.xlsx",
        }
        dynamic.setdefault("tacs_safety_adverse_events", {})["missing_data_note_counts_project_evidence"] = {
            "value": clinical_summary["missing_note_counts"],
            "evidence": "F:/CJZFile/EEG_M1/脑卒中患者信息记录表.xlsx",
        }
    cnt_header = collect_cnt_header_boundary()
    if cnt_header:
        dynamic.setdefault("eeg_hardware_reference_impedance", {})["raw_cnt_header_boundary_project_evidence"] = {
            "value": cnt_header,
            "evidence": "F:/CJZFile/EEG_M1/Patient_tACS_M1_EEG representative raw CNT headers",
        }
        dynamic.setdefault("raw_eeg_preprocessing", {})["raw_cnt_header_boundary_project_evidence"] = {
            "value": cnt_header,
            "evidence": "F:/CJZFile/EEG_M1/Patient_tACS_M1_EEG representative raw CNT headers",
        }
    pdf_audit_summary = collect_patient_record_pdf_audit_summary()
    if pdf_audit_summary:
        for field_id in ("ethics_approval", "informed_consent", "tacs_safety_adverse_events"):
            dynamic.setdefault(field_id, {})["patient_record_pdf_text_audit_project_evidence"] = {
                "value": pdf_audit_summary,
                "evidence": "docs/patient_record_pdf_text_audit.md; results/tables/patient_record_pdf_text_audit.csv",
            }
    return dynamic


def collect_patient_cnt_window() -> str:
    if not PATIENT_RAW_CNT_DIR.exists():
        return ""
    observations: list[tuple[datetime, str]] = []
    phase_counts: Counter[str] = Counter()
    versions: Counter[str] = Counter()
    for path in PATIENT_RAW_CNT_DIR.rglob("*.cnt"):
        rel_parts = path.relative_to(PATIENT_RAW_CNT_DIR).parts
        phase_counts[rel_parts[0] if rel_parts else "unknown"] += 1
        header = read_ascii_header(path, 256)
        version = first_match(header, r"Version\s+([0-9.]+)")
        if version:
            versions[version] += 1
        dt = parse_cnt_header_datetime(header)
        if dt is None:
            dt = datetime.fromtimestamp(path.stat().st_mtime)
        observations.append((dt, rel_parts[0] if rel_parts else "unknown"))
    if not observations:
        return ""
    dates = [item[0] for item in observations]
    phase_text = ", ".join(f"{phase} {count}" for phase, count in sorted(phase_counts.items()))
    version_text = ", ".join(f"Version {version} ({count})" for version, count in sorted(versions.items())) or "version not parsed"
    return (
        f"Raw patient CNT files currently available under Patient_tACS_M1_EEG span "
        f"{min(dates).date().isoformat()} to {max(dates).date().isoformat()} across {len(observations)} files "
        f"({phase_text}); parsed CNT header versions: {version_text}. This is an EEG recording-file window, "
        "not a verified recruitment window, and requires author confirmation before use as study-date wording."
    )


def collect_clinical_workbook_summary() -> dict[str, str]:
    if not PATIENT_CLINICAL_WORKBOOK.exists():
        return {}
    wb = load_workbook(PATIENT_CLINICAL_WORKBOOK, data_only=True, read_only=True)
    ws = wb.active
    header_row = None
    headers: list[str] = []
    for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        values = [clean_cell(value) for value in row]
        if values and values[0] == "编号" and "脱落原因" in values:
            header_row = row_idx
            headers = values
            break
    if header_row is None:
        return {}
    missing_idx = headers.index("缺少数据") if "缺少数据" in headers else None
    reason_idx = headers.index("脱落原因") if "脱落原因" in headers else None
    m1_rows = 0
    missing_notes: list[str] = []
    reason_notes: list[str] = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        values = [clean_cell(value) for value in row]
        if not values or not values[0]:
            continue
        if not str(values[0]).lower().startswith("sub"):
            break
        m1_rows += 1
        if missing_idx is not None and missing_idx < len(values) and values[missing_idx]:
            missing_notes.append(values[missing_idx])
        if reason_idx is not None and reason_idx < len(values) and values[reason_idx]:
            reason_notes.append(values[reason_idx])

    record_headers = []
    if CLINICAL_RECORD_WORKBOOK.exists():
        record_wb = load_workbook(CLINICAL_RECORD_WORKBOOK, data_only=True, read_only=True)
        record_ws = record_wb.active
        for row in record_ws.iter_rows(values_only=True):
            values = [clean_cell(value) for value in row]
            if values and values[0] == "编号":
                record_headers = [value for value in values if value]
                break
    scale_headers = [header for header in sorted(set(headers + record_headers)) if any(term in header for term in ("FMA", "MBI", "BBT", "MMSE"))]
    clinical_fields = (
        "The M1 patient-information workbook contains non-identifying clinical columns for subject ID, age, disease duration, sex, "
        "affected side, pre/post FMA, pre/post MBI, missing-data notes, dropout reason, and MRI count. A second M1 clinical "
        "record workbook also contains BBT and MMSE columns. These source fields support descriptive cohort and timing summaries, "
        "but they do not define inclusion criteria, exclusion criteria, or stroke-subtype rules."
    )
    scale_fields = (
        f"Clinical source workbooks contain scale fields including {', '.join(scale_headers) if scale_headers else 'FMA/MBI fields'}. "
        "They do not identify the assessor credentials or blinding status."
    )
    category_counts = categorize_reason_notes(reason_notes)
    category_text = ", ".join(f"{name}: {count}" for name, count in category_counts.items()) or "no reason categories parsed"
    missing_note_counts = (
        f"The M1 patient-information workbook contains {m1_rows} M1 rows; {len(missing_notes)} rows have nonblank missing-data notes "
        f"and {len(reason_notes)} rows have nonblank dropout/reason notes. Non-identifying reason categories parsed from the source "
        f"notes are: {category_text}. This supports participant-flow review only and is not a complete adverse-event monitoring statement."
    )
    return {
        "clinical_fields": clinical_fields,
        "scale_fields": scale_fields,
        "missing_note_counts": missing_note_counts,
    }


def collect_cnt_header_boundary() -> str:
    if not PATIENT_RAW_CNT_DIR.exists():
        return ""
    labels_seen: set[str] = set()
    versions: set[str] = set()
    for path in sorted(PATIENT_RAW_CNT_DIR.rglob("*.cnt"))[:12]:
        header = read_ascii_header(path, 4096)
        version = first_match(header, r"Version\s+([0-9.]+)")
        if version:
            versions.add(version)
        for label in ("FP1", "FPZ", "FP2", "AF3", "AF4", "FC3", "FCZ", "CP3", "CPZ", "OZ"):
            if label in header.upper():
                labels_seen.add(label)
    if not labels_seen and not versions:
        return ""
    return (
        f"Representative raw CNT headers expose {', '.join('Version ' + item for item in sorted(versions)) or 'a CNT version string'} "
        f"and extended 10-20/10-10 style labels including {', '.join(sorted(labels_seen)) if labels_seen else 'standard EEG labels'}. "
        "The inspected headers do not provide a submission-ready amplifier model, cap system, online reference, ground, impedance threshold, "
        "or upstream artifact-preprocessing protocol."
    )


def collect_patient_record_pdf_audit_summary() -> str:
    if not PATIENT_PDF_AUDIT_CSV.exists():
        return ""
    with PATIENT_PDF_AUDIT_CSV.open("r", encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    if not rows:
        return ""
    status_counts = Counter(row.get("text_layer_status", "unknown") for row in rows)
    utility_counts = Counter(row.get("metadata_utility", "unknown") for row in rows)
    total_pages = sum(int(row.get("page_count") or 0) for row in rows)
    total_text = sum(int(row.get("extracted_text_chars") or 0) for row in rows)
    status_text = ", ".join(f"{key} {value}" for key, value in sorted(status_counts.items()))
    utility_text = ", ".join(f"{key} {value}" for key, value in sorted(utility_counts.items()))
    return (
        f"A non-identifying audit of local patient/healthy record-book PDFs inspected {len(rows)} PDFs and {total_pages} pages. "
        f"Extractable text was minimal ({total_text} compact characters total); text-layer status counts were {status_text}; "
        f"metadata-utility counts were {utility_text}. These scanned records cannot support automatic ethics, consent, safety, "
        "or protocol prefill without OCR or manual review."
    )


def read_ascii_header(path: Path, byte_count: int) -> str:
    try:
        with path.open("rb") as handle:
            data = handle.read(byte_count)
    except OSError:
        return ""
    text = data.decode("latin1", errors="ignore")
    return " ".join(re.findall(r"[ -~]{2,}", text))


def parse_cnt_header_datetime(header: str) -> datetime | None:
    match = re.search(r"(\d{2})/(\d{2})/(\d{2})\s+(\d{2}:\d{2}:\d{2})", header)
    if not match:
        return None
    day, month, year, time_text = match.groups()
    return datetime.strptime(f"20{year}-{month}-{day} {time_text}", "%Y-%m-%d %H:%M:%S")


def first_match(text: str, pattern: str) -> str:
    match = re.search(pattern, text, flags=re.I)
    return match.group(1) if match else ""


def clean_cell(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def categorize_reason_notes(notes: list[str]) -> Counter[str]:
    categories: Counter[str] = Counter()
    for note in notes:
        if "未接受治疗" in note:
            categories["non-treatment"] += 1
        elif "出院" in note:
            categories["discharge"] += 1
        elif any(term in note for term in ("依从", "认知", "交流")):
            categories["poor compliance or cognitive/communication difficulty"] += 1
        elif "热" in note:
            categories["EEG-cap heat discomfort"] += 1
        elif "高血压" in note:
            categories["post-session discomfort followed by hypertension"] += 1
        elif "核磁" in note:
            categories["MRI-related discomfort"] += 1
        elif "不适" in note or "难受" in note:
            categories["post-session discomfort"] += 1
        else:
            categories["other"] += 1
    return categories


def build_report(rows: list[dict[str, str]]) -> str:
    lines = [
        "# Author Metadata Project-Evidence Prefill Report",
        "",
        "This report lists metadata values that can be prefilled from current project evidence. It is not an author-approved final metadata file. Ethics, consent, registry, data-sharing permission, repository DOI/licence, hardware details, safety summaries, and authorship declarations still require author or institutional confirmation.",
        "",
        "## Prefilled Values",
        "",
        "| Field | Metadata key | Prefilled value | Evidence |",
        "|---|---|---|---|",
    ]
    for row in rows:
        lines.append(
            f"| {escape(row['field_id'])} | {escape(row['metadata_key'])} | {escape(row['value'])} | {escape(row['evidence'])} |"
        )
    lines.extend(
        [
            "",
            "## Source Discovery Notes",
            "",
            "- The patient-information workbook contains two table blocks: an M1 patient group and a healthy-control group. Only the M1 patient block was used for disease-duration and participant-flow evidence.",
            "- The M1 block contains 29 patient entries; the final supervised cohort is the 19-patient list in `F:/CJZFile/EEG_M1/19例患者脑电数据完整性检查.xlsx`.",
            "- Nine M1 entries contain missing-data or discontinuation notes. These notes can support participant-flow review, but they do not establish a complete adverse-event collection method, severity grading, causality adjudication, or denominator for formal safety reporting.",
            "- Raw patient CNT headers and file inventory can support an EEG recording-file window, but not a verified recruitment window or final study-design statement.",
            "- Representative raw CNT headers expose version/date/channel-label evidence, but not amplifier, cap, reference, ground, impedance, or upstream artifact-preprocessing details.",
            "- Patient and healthy record-book PDFs were audited with a non-identifying text-layer check. They are image-based scans and require OCR or manual review before any ethics, consent, safety, or protocol wording can be inferred.",
            "- No ethics approval number, consent wording, study registration record, recruitment date window, EEG amplifier/reference/ground/impedance protocol, conventional-rehabilitation protocol, repository DOI/licence, or final authorship declarations were located in the source files reviewed for this prefill.",
        ]
    )
    lines.extend(
        [
            "",
            "## Still Author-Required",
            "",
        ]
    )
    for field_id in AUTHOR_REQUIRED_FIELDS:
        lines.append(f"- `{field_id}`")
    lines.extend(
        [
            "",
            "## Use",
            "",
            "1. Review `outputs/manuscript_package/author_submission_metadata_project_prefill.json`.",
            "2. Copy only author-approved values into `docs/author_submission_metadata_template.json` or the XLSX intake workbook.",
            "3. Complete all remaining author-required fields.",
            "4. Run strict validation and the final insertion protocol before editing clean manuscript files.",
            "",
        ]
    )
    return "\n".join(lines)


def escape(value: Any) -> str:
    return str(value).replace("|", "\\|").replace("\n", " ")


if __name__ == "__main__":
    main()
