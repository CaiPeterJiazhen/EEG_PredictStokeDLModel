from __future__ import annotations

import re
import os
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
INPUT_MD = ROOT / "docs" / "manuscript_residual_aware_ssl_cnn_nature_clean_zh.md"
OUTPUT_DOCX = Path(
    os.environ.get(
        "OUTPUT_DOCX",
        ROOT / "output" / "doc" / "ResidualAware_SSL_CNN_Nature_Clean_ZH.docx",
    )
)
EMBED_FIGURES = os.environ.get("CHINESE_DRAFT_EMBED_FIGURES", "1") != "0"

FIGURES = {
    "图 1": ROOT / "results" / "figures" / "nature" / "figure1_study_design_model.png",
    "图 2": ROOT / "results" / "figures" / "nature" / "figure2_performance_calibration.png",
    "图 3": ROOT / "results" / "figures" / "nature" / "figure3_robustness_ablation.png",
    "图 4": ROOT / "results" / "figures" / "nature" / "figure4_explainability_neurophysiology.png",
    "补充图 1": ROOT / "results" / "figures" / "nature" / "supplementary_error_subjects.png",
    "补充图 2": ROOT
    / "results"
    / "figures"
    / "explainability"
    / "mne_wpli_connectivity"
    / "mne_wpli_connectivity_contact_sheet.png",
    "补充图 3": ROOT / "results" / "figures" / "nature" / "supplementary_performance_precision.png",
    "补充图 4": ROOT / "results" / "figures" / "nature" / "supplementary_clinical_incremental_value.png",
}


def set_run_font(run, *, size: float = 10.5, bold: bool = False, mono: bool = False) -> None:
    run.bold = bold
    run.font.size = Pt(size)
    if mono:
        run.font.name = "Consolas"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    else:
        run.font.name = "Times New Roman"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    for style_name, size, east_asia in [
        ("Normal", 10.5, "宋体"),
        ("Title", 18, "黑体"),
        ("Heading 1", 15, "黑体"),
        ("Heading 2", 12, "黑体"),
        ("Heading 3", 10.5, "黑体"),
    ]:
        style = document.styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    document.styles["Normal"].paragraph_format.line_spacing = 1.15
    document.styles["Normal"].paragraph_format.space_after = Pt(4)


def add_inline_runs(paragraph, text: str) -> None:
    parts = re.split(r"(`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, size=9, mono=True)
        else:
            run = paragraph.add_run(part)
            set_run_font(run)


def add_picture_if_needed(document: Document, line: str) -> None:
    if not EMBED_FIGURES:
        return
    for label, path in FIGURES.items():
        if line.startswith(f"{label}.") and path.exists():
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run()
            width = 6.4 if label.startswith("图") else 6.8
            run.add_picture(str(path), width=Inches(width))
            return


def build_docx() -> None:
    OUTPUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    configure_document(document)

    in_references = False
    for raw_line in INPUT_MD.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith("# "):
            title = document.add_paragraph(style="Title")
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = title.add_run(line[2:])
            set_run_font(run, size=18, bold=True)
            continue
        if line.startswith("## "):
            heading = line[3:]
            in_references = heading == "参考文献"
            document.add_heading(heading, level=1)
            continue
        if line.startswith("### "):
            document.add_heading(line[4:], level=2)
            continue

        add_picture_if_needed(document, line)

        paragraph = document.add_paragraph()
        if in_references and re.match(r"^\d+\. ", line):
            paragraph.style = document.styles["List Number"]
            line = re.sub(r"^\d+\. ", "", line)
        add_inline_runs(paragraph, line)

    document.save(OUTPUT_DOCX)
    print(OUTPUT_DOCX)


if __name__ == "__main__":
    build_docx()
