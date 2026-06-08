from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.shared import RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
SOURCE_MD = Path(r"C:\Users\HPGZZ\Downloads\manuscript_optimized_expanded_zh.md")
OUT_MD = Path(r"C:\Users\HPGZZ\Downloads\manuscript_optimized_expanded_zh_with_figures.md")
OUT_DOCX = Path(r"C:\Users\HPGZZ\Downloads\manuscript_optimized_expanded_zh_with_figures.docx")
FIG_ROOT = PROJECT_ROOT / "results" / "figures" / "revised_initial"

FIGURES: dict[str, list[tuple[str, Path]]] = {
    "图1": [("图1 总体技术框架", FIG_ROOT / "figure1_overall_framework_provided.png")],
    "图2": [("图2 患者层面 Barlow Twins SSL 预训练", FIG_ROOT / "figure2_ssl_framework_provided.png")],
    "图3": [("图3 残差感知 SSL-CNN 结构", FIG_ROOT / "figure3_cnn_residual_aware_provided.png")],
    "图4": [
        ("图4A 最终模型 ROC 曲线", FIG_ROOT / "figure4a_final_model_roc.png"),
        ("图4B 最终模型混淆矩阵", FIG_ROOT / "figure4b_final_model_confusion_matrix.png"),
        ("图4C-a 不同模型患者层面指标直方图", FIG_ROOT / "figure4c_a_model_metric_histogram.png"),
        ("图4C-b 不同模型 Brier 校准误差", FIG_ROOT / "figure4c_b_brier_calibration.png"),
        ("图4D 最终模型训练和验证损失曲线", FIG_ROOT / "figure4d_final_model_loss_curves.png"),
    ],
    "图5": [("图5 稳定性与消融", FIG_ROOT / "figure5_stability_ablation.png")],
    "图6": [
        ("图6A PSD 频段归因 topomap", FIG_ROOT / "figure6a_psd_topomap_bands.png"),
        ("图6B WPLI/FC 频段连接归因图", FIG_ROOT / "figure6b_wpli_connectivity_bands.png"),
    ],
}

TABLE5_REPLACEMENT = """| PSD | EO | Gamma | TP7 | -0.000892 | 0.001330 |
| PSD | EO | Gamma | C5 | -0.000846 | 0.001245 |
| PSD | EO | Beta High | TP7 | -0.000749 | 0.001228 |
| PSD | EO | Beta High | FPZ | -0.000943 | 0.001159 |
| PSD | EO | Delta | F5 | -0.000817 | 0.001137 |"""


def main() -> None:
    text = SOURCE_MD.read_text(encoding="utf-8-sig")
    enhanced = _insert_figures_and_update_values(text)
    OUT_MD.write_text(enhanced, encoding="utf-8")
    build_docx(enhanced, OUT_DOCX)
    print(f"Wrote Markdown: {OUT_MD}")
    print(f"Wrote DOCX: {OUT_DOCX}")


def _insert_figures_and_update_values(text: str) -> str:
    caption_replacements = {
        "**图1｜总体技术框架。** 建议放置整体流程图：左侧为基线 EEG 和临床量表采集；中部为预处理、患侧对齐、PSD/WPLI 提取、Barlow SSL 预训练和残差感知监督训练；右侧为 LOSO 评估、统计验证和可解释性输出。图中应明确标注：治疗后 FMA-UE 只用于标签构建，模型输入限定为治疗前 EEG。":
            "**图1｜总体技术框架。** 基线 EEG 与临床量表采集、EEG 预处理、患侧对齐、PSD/WPLI 特征提取、患者层面 Barlow SSL 预训练、残差感知监督训练、LOSO 评估和可解释性分析流程。治疗后 FMA-UE 仅用于标签构建，模型输入限定为治疗前 EEG。",
        "**图2｜患者层面 Barlow Twins SSL 预训练。** 建议展示两个增强视图、共享编码器、投影头、交叉相关矩阵和冗余降低损失，并标注每个 LOSO 折中测试患者被排除出 SSL 池。":
            "**图2｜患者层面 Barlow Twins SSL 预训练。** 对同一患者 EEG 特征生成两个增强视图，输入共享编码器和投影头，通过 Barlow Twins 交叉相关矩阵实现不变性学习和冗余降低。每个 LOSO 折中，测试患者在缩放、增强和 SSL 编码器训练之前均被排除出训练池。",
        "**图3｜残差感知 SSL-CNN 结构。** 建议展示 PSD/WPLI 分支、EO/EC gated fusion、患者层面 embedding、二分类头、残差回归头和排序/软标签训练目标。图中应突出：训练阶段多目标，推断阶段单一二分类概率。":
            "**图3｜残差感知 SSL-CNN 结构。** 模型包含 PSD 与 WPLI 分支、EO/EC 门控融合、患者层面 EEG embedding、二分类比例恢复头，以及训练阶段使用的残差回归、成对排序和残差距离软标签辅助目标。推断阶段仅使用二分类头输出比例恢复概率。",
        "**图4｜模型性能和训练过程。** 建议放置 4 个子图：A，最终模型 ROC 曲线；B，最终模型混淆矩阵；C，传统 ML、no-SSL CNN 和 residual-aware SSL-CNN 的患者层面指标热图；D，最终模型训练和验证损失曲线。正文应强调：图4C 中 locked subject-level summary 用于展示分类、排序和校准表现；图5 的 10 种子分布用于展示训练稳定性。":
            "**图4｜模型性能和训练过程。** A，最终残差感知 SSL-CNN 的锁定患者层面 ROC 曲线；B，固定阈值 0.5 下的患者层面混淆矩阵；C-a，传统 EEG-ML、no-SSL CNN、Barlow CNN 和残差感知 SSL-CNN 的分类与排序指标直方图；C-b，各模型 Brier 校准误差对比；D，最终模型训练和验证总损失、分类损失及加权残差感知辅助损失曲线。",
        "**图5｜稳定性与消融。** 建议放置种子级 accuracy 分布、核心 ablation 热图、特征/状态/频段消融排名和输入维度-ROC-AUC 信息效率散点图。图中应以 residual-aware SSL-CNN 的平均 accuracy、最小 accuracy 和 SD 为主要可视化重点，而不突出 no-SSL CNN 的单次或 ensemble accuracy。":
            "**图5｜稳定性与消融。** 展示核心深度模型的 10 种子 accuracy 分布、SSL/残差感知核心消融指标热图、特征/状态/频段消融排名，以及输入维度与 ROC-AUC/PR-AUC 的信息效率关系。",
        "**图6｜EEG 解释性结果。** 建议分为两个主图：A，PSD 频段 topomap，展示 EO/EC 条件下 delta 至 gamma 七个频段的 signed attribution；B，WPLI connectivity 图，展示 EO/EC 条件下每个频段 top-20 连接边，边颜色表示归因方向，线宽表示绝对归因强度。图注应强调这些图用于生成候选 EEG 生物标志物假设，而非确认因果机制。":
            "**图6｜EEG 解释性结果。** A，PSD 频段 topomap 展示 EO/EC 条件下 delta 至 gamma 七个频段的 signed attribution；B，WPLI connectivity 图展示 EO/EC 条件下各频段 top-20 连接边，边颜色表示归因方向，线宽表示绝对归因强度。这些图用于生成候选 EEG 生物标志物假设，而非确认因果机制。",
    }
    for old, new in caption_replacements.items():
        text = text.replace(old, new)
    text = text.replace(
        "| PSD | EO | Other | TP7 | -0.000874 | 0.001308 |\n"
        "| PSD | EO | Delta | F5 | -0.000929 | 0.001245 |\n"
        "| PSD | EO | Beta High | TP7 | -0.000749 | 0.001228 |\n"
        "| PSD | EO | Other | C5 | -0.000820 | 0.001211 |\n"
        "| PSD | EO | Beta High | FPZ | -0.000943 | 0.001159 |",
        TABLE5_REPLACEMENT,
    )
    text = text.replace(
        "PSD 归因集中在 EO 条件下的额叶、颞部和围中央通道，代表性通道包括 F5、T7、TP7、FC5、C5 和 FPZ。频段上，delta 和 beta-high 特征均进入 top attribution。",
        "PSD 归因集中在 EO 条件下的额叶、颞部和围中央通道，代表性通道包括 F5、T7、TP7、FC5、C5 和 FPZ。频段上，gamma、beta-high 和 delta 特征均进入 top attribution。",
    )
    text = text.replace(
        "A，PSD 频段 topomap，展示 EO/EC 条件下六个频段的 signed attribution；B，WPLI connectivity 图",
        "A，PSD 频段 topomap，展示 EO/EC 条件下 delta 至 gamma 七个频段的 signed attribution；B，WPLI connectivity 图",
    )
    text = text.replace(
        "**图6｜EEG 解释性结果。** 建议分为两个主图：A，PSD 频段 topomap，展示 EO/EC 条件下 delta 至 gamma 七个频段的 signed attribution；B，WPLI connectivity 图，展示 EO/EC 条件下每个频段 top-20 连接边，边颜色表示归因方向，线宽表示绝对归因强度。图注应强调这些图用于生成候选 EEG 生物标志物假设，而非确认因果机制。",
        "**图6｜EEG 解释性结果。** A，PSD 频段 topomap 展示 EO/EC 条件下 delta 至 gamma 七个频段的 signed attribution；B，WPLI connectivity 图展示 EO/EC 条件下各频段 top-20 连接边，边颜色表示归因方向，线宽表示绝对归因强度。这些图用于生成候选 EEG 生物标志物假设，而非确认因果机制。",
    )

    output: list[str] = []
    for line in text.splitlines():
        output.append(line)
        match = re.match(r"^\*\*(图[1-6])｜", line.strip())
        if match:
            key = match.group(1)
            for alt, path in FIGURES.get(key, []):
                if not path.exists():
                    raise FileNotFoundError(path)
                output.append("")
                output.append(f"![{alt}]({path.as_posix()})")
    return "\n".join(output).rstrip() + "\n"


def build_docx(markdown: str, output: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    _configure_styles(doc)

    lines = markdown.splitlines()
    i = 0
    first_title = True
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        if line.strip() == "---":
            doc.add_page_break()
            i += 1
            continue
        if line.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            rows = _parse_markdown_table(table_lines)
            if rows:
                _add_three_line_table(doc, rows)
            continue
        image_match = re.match(r"^!\[(.*?)\]\((.*?)\)$", line.strip())
        if image_match:
            _add_image(doc, Path(image_match.group(2)), image_match.group(1))
            i += 1
            continue
        if line.startswith("\\["):
            math_lines = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("\\]"):
                if lines[i].strip():
                    math_lines.append(lines[i].strip())
                i += 1
            if i < len(lines) and lines[i].strip().startswith("\\]"):
                i += 1
            _add_formula(doc, " ".join(math_lines))
            continue
        if line.startswith("# "):
            paragraph = doc.add_paragraph()
            paragraph.style = "ManuscriptTitle"
            _add_inline_runs(paragraph, line[2:].strip(), bold_default=False)
            first_title = False
            i += 1
            continue
        if line.startswith("## "):
            paragraph = doc.add_paragraph(style="Heading 1")
            _add_inline_runs(paragraph, line[3:].strip(), bold_default=False)
            i += 1
            continue
        if line.startswith("### "):
            paragraph = doc.add_paragraph(style="Heading 2")
            _add_inline_runs(paragraph, line[4:].strip(), bold_default=False)
            i += 1
            continue
        if re.match(r"^\d+\.\s", line.strip()):
            paragraph = doc.add_paragraph(style="Reference")
            _add_inline_runs(paragraph, line.strip(), bold_default=False)
            i += 1
            continue

        style = "Caption" if line.startswith("**图") or line.startswith("**表") else "Normal"
        paragraph = doc.add_paragraph(style=style)
        _add_inline_runs(paragraph, line.strip(), bold_default=False)
        if first_title:
            first_title = False
        i += 1

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def _configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "SimSun"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)

    title = styles["ManuscriptTitle"] if "ManuscriptTitle" in [style.name for style in styles] else styles.add_style("ManuscriptTitle", 1)
    title.font.name = "SimHei"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    title.font.size = Pt(16)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(12)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    h1 = styles["Heading 1"]
    h1.font.name = "SimHei"
    h1._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    h1.font.size = Pt(14)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor(0, 0, 0)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)

    h2 = styles["Heading 2"]
    h2.font.name = "SimHei"
    h2._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    h2.font.size = Pt(12)
    h2.font.bold = True
    h2.font.color.rgb = RGBColor(0, 0, 0)
    h2.paragraph_format.space_before = Pt(10)
    h2.paragraph_format.space_after = Pt(4)

    caption = styles["Caption"] if "Caption" in [style.name for style in styles] else styles.add_style("Caption", 1)
    caption.font.name = "SimSun"
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    caption.font.size = Pt(9)
    caption.font.bold = False
    caption.font.color.rgb = RGBColor(0, 0, 0)
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(4)
    caption.paragraph_format.keep_with_next = True

    reference = styles["Reference"] if "Reference" in [style.name for style in styles] else styles.add_style("Reference", 1)
    reference.font.name = "Times New Roman"
    reference._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    reference.font.size = Pt(9)
    reference.paragraph_format.first_line_indent = Inches(-0.22)
    reference.paragraph_format.left_indent = Inches(0.22)
    reference.paragraph_format.space_after = Pt(3)


def _add_inline_runs(paragraph, text: str, *, bold_default: bool) -> None:
    text = re.sub(r"\\\((.*?)\\\)", lambda match: _latex_to_readable(match.group(1)), text)
    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        if not part:
            continue
        bold = bold_default
        if part.startswith("**") and part.endswith("**"):
            part = part[2:-2]
            bold = True
        run = paragraph.add_run(part)
        run.bold = bold
        run.font.name = "Times New Roman"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def _add_formula(doc: Document, latex: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(_latex_to_readable(latex))
    run.font.name = "Cambria Math"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Cambria Math")
    run.font.size = Pt(10.5)


def _latex_to_readable(value: str) -> str:
    text = value.strip()
    replacements = {
        r"\Delta": "Δ",
        r"\times": "×",
        r"\quad": "  ",
        r"\alpha": "α",
        r"\beta": "β",
        r"\gamma": "γ",
        r"\lambda": "λ",
        r"\sum": "Σ",
        r"\ne": "≠",
        r"\le": "≤",
        r"\ge": "≥",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    text = re.sub(r"\\mathcal\{L\}_\{([^}]+)\}", r"L_\1", text)
    text = re.sub(r"\\mathcal\{L\}", "L", text)
    text = text.replace("\\", "")
    text = text.replace("{", "").replace("}", "")
    text = re.sub(r"\s+", " ", text)
    return text


def _parse_markdown_table(lines: list[str]) -> list[list[str]]:
    rows: list[list[str]] = []
    for line in lines:
        stripped = line.strip()
        cells = [cell.strip() for cell in stripped.strip("|").split("|")]
        if all(re.fullmatch(r":?-{2,}:?", cell.replace(" ", "")) for cell in cells):
            continue
        rows.append(cells)
    return rows


def _add_three_line_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    n_cols = max(len(row) for row in rows)
    rows = [row + [""] * (n_cols - len(row)) for row in rows]
    table = doc.add_table(rows=len(rows), cols=n_cols)
    table.autofit = False
    _set_table_width(table, 9360)
    widths = _column_widths(rows, 9360)
    for col_index, width in enumerate(widths):
        for cell in table.columns[col_index].cells:
            _set_cell_width(cell, width)
    font_size = Pt(8 if n_cols >= 7 else 9)
    for row_index, row in enumerate(rows):
        for col_index, value in enumerate(row):
            cell = table.cell(row_index, col_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cell, top=80, bottom=80, start=120, end=120)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if row_index == 0 or _is_short_or_numeric(value) else WD_ALIGN_PARAGRAPH.LEFT
            run = paragraph.add_run(value)
            run.font.size = font_size
            run.font.name = "Times New Roman"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
            if row_index == 0:
                run.bold = True
    _apply_three_line_borders(table)
    doc.add_paragraph()


def _column_widths(rows: list[list[str]], total_dxa: int) -> list[int]:
    n_cols = len(rows[0])
    if n_cols == 2:
        return [int(total_dxa * 0.32), int(total_dxa * 0.68)]
    weights = []
    for col in range(n_cols):
        values = [row[col] for row in rows]
        max_len = max(len(value) for value in values)
        numeric_ratio = sum(_is_numeric_like(value) for value in values[1:]) / max(len(values) - 1, 1)
        weight = max_len * (0.75 if numeric_ratio > 0.7 else 1.0)
        weights.append(max(weight, 8))
    total_weight = sum(weights)
    raw = [int(total_dxa * weight / total_weight) for weight in weights]
    min_width = 720 if n_cols >= 7 else 900
    raw = [max(width, min_width) for width in raw]
    scale = total_dxa / sum(raw)
    return [int(width * scale) for width in raw]


def _apply_three_line_borders(table) -> None:
    for row in table.rows:
        for cell in row.cells:
            _set_cell_borders(cell, top=None, bottom=None, left=None, right=None)
    for cell in table.rows[0].cells:
        _set_cell_borders(cell, top={"val": "single", "sz": "12"}, bottom={"val": "single", "sz": "8"})
    for cell in table.rows[-1].cells:
        _set_cell_borders(cell, bottom={"val": "single", "sz": "12"})


def _set_cell_borders(cell, **borders) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = tc_borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            tc_borders.append(element)
        spec = borders.get(edge)
        if spec is None:
            element.set(qn("w:val"), "nil")
        else:
            element.set(qn("w:val"), spec.get("val", "single"))
            element.set(qn("w:sz"), spec.get("sz", "8"))
            element.set(qn("w:color"), spec.get("color", "000000"))


def _set_table_width(table, width_dxa: int) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(width_dxa))
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), "0")
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_ind)


def _set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.tcW
    tc_w.type = "dxa"
    tc_w.w = width_dxa


def _set_cell_margins(cell, top: int, bottom: int, start: int, end: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in {"top": top, "bottom": bottom, "start": start, "end": end}.items():
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _add_image(doc: Document, path: Path, alt: str) -> None:
    if not path.exists():
        raise FileNotFoundError(path)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run()
    width = Inches(6.25)
    if any(key in path.name for key in ("figure4a", "figure4b")):
        width = Inches(4.8)
    run.add_picture(str(path), width=width)
    paragraph.paragraph_format.keep_together = True


def _is_numeric_like(value: str) -> bool:
    value = value.strip()
    return bool(re.fullmatch(r"(NA|<)?-?\d+(\.\d+)?(\s*±\s*\d+(\.\d+)?)?|\d+/\d+|[-–]|\d+%", value))


def _is_short_or_numeric(value: str) -> bool:
    return _is_numeric_like(value) or len(value) <= 18


if __name__ == "__main__":
    main()
