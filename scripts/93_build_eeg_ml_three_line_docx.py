from __future__ import annotations

from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


PROJECT_ROOT = Path(__file__).resolve().parents[1]
SOURCE_TABLE = PROJECT_ROOT / "results" / "tables" / "eeg_only_ml_three_line_table.csv"
OUTPUT_DOC = PROJECT_ROOT / "output" / "doc" / "eeg_only_ml_three_line_table.docx"


def set_run_font(run, size_pt: float, bold: bool = False) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r_fonts.set(qn("w:ascii"), "Times New Roman")
    r_fonts.set(qn("w:hAnsi"), "Times New Roman")


def set_cell_text(cell, text: str, size_pt: float = 8.5, bold: bool = False, align=WD_ALIGN_PARAGRAPH.CENTER) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(str(text))
    set_run_font(run, size_pt=size_pt, bold=bold)


def clear_cell_text(cell) -> None:
    paragraph = cell.paragraphs[0]
    for run in paragraph.runs:
        run._element.getparent().remove(run._element)


def set_cell_borders(cell, top=None, bottom=None, left=None, right=None) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)

    def border_element(edge: str, spec) -> None:
        tag = f"w:{edge}"
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        if spec is None:
            element.set(qn("w:val"), "nil")
            return
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), str(spec.get("sz", 8)))
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), spec.get("color", "000000"))

    border_element("top", top)
    border_element("bottom", bottom)
    border_element("left", left)
    border_element("right", right)


def set_table_width(table, width_cm: float) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(int(width_cm * 567)))
    tbl_w.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_cm: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def build_document(table_data: pd.DataFrame) -> Document:
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.3)
    section.left_margin = Cm(1.2)
    section.right_margin = Cm(1.2)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("EEG-only machine-learning performance")
    set_run_font(run, size_pt=14, bold=True)
    title.paragraph_format.space_after = Pt(8)

    columns = list(table_data.columns)
    table = document.add_table(rows=len(table_data) + 1, cols=len(columns))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_width(table, 27.2)

    col_widths = [6.4, 2.6, 3.6, 2.6, 2.6, 2.5, 2.5, 2.4]
    top_line = {"sz": 12, "color": "000000"}
    mid_line = {"sz": 8, "color": "000000"}
    bottom_line = {"sz": 12, "color": "000000"}

    for col_idx, column in enumerate(columns):
        cell = table.cell(0, col_idx)
        clear_cell_text(cell)
        set_cell_width(cell, col_widths[col_idx])
        set_cell_text(
            cell,
            column,
            size_pt=8.8,
            bold=True,
            align=WD_ALIGN_PARAGRAPH.LEFT if col_idx == 0 else WD_ALIGN_PARAGRAPH.CENTER,
        )
        set_cell_borders(cell, top=top_line, bottom=mid_line)

    for row_idx, (_, row) in enumerate(table_data.iterrows(), start=1):
        for col_idx, column in enumerate(columns):
            cell = table.cell(row_idx, col_idx)
            clear_cell_text(cell)
            set_cell_width(cell, col_widths[col_idx])
            set_cell_text(
                cell,
                row[column],
                size_pt=8.4,
                align=WD_ALIGN_PARAGRAPH.LEFT if col_idx == 0 else WD_ALIGN_PARAGRAPH.CENTER,
            )
            set_cell_borders(cell)

    last_row = table.rows[-1]
    for cell in last_row.cells:
        set_cell_borders(cell, bottom=bottom_line)

    note = document.add_paragraph()
    note.paragraph_format.space_before = Pt(8)
    note.paragraph_format.space_after = Pt(0)
    note.alignment = WD_ALIGN_PARAGRAPH.LEFT
    note_run = note.add_run(
        "注：所有结果均为 19 例监督队列的 patient-level LOSO 点估计；"
        "Brier 越低越好，其余指标越高越好。"
    )
    set_run_font(note_run, size_pt=8.5)
    return document


def main() -> None:
    OUTPUT_DOC.parent.mkdir(parents=True, exist_ok=True)
    table_data = pd.read_csv(SOURCE_TABLE)
    document = build_document(table_data)
    document.save(OUTPUT_DOC)
    print(f"Wrote {OUTPUT_DOC}")


if __name__ == "__main__":
    main()
